"""Strict historical Aptitude answer import; never trusts spreadsheet points."""
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from datetime import datetime
from openpyxl import load_workbook
from domain import clean_text, grade, require, student_key


def normalized(value):
    value = unicodedata.normalize('NFKC', str(value or '')).casefold()
    return re.sub(r'[^\w%]+', '', value)


def read_responses(raw, questions):
    require(questions and all(q['type'] in {'single', 'boolean'} and q.get('key') for q in questions),
            'Set an answer key for every Aptitude multiple-choice question before importing.')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            require(sum(i.file_size for i in archive.infolist()) <= 50 * 1024 * 1024, 'Workbook is too large when expanded.')
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        require(False, 'Choose a valid .xlsx workbook.')
    try:
        require('Input Data' in book.sheetnames, 'The workbook must contain the Input Data sheet from the score tracker.')
        sheet = book['Input Data']
        require(sheet.max_row <= 5001 and sheet.max_column <= 200, 'Import supports up to 5,000 rows and 200 columns.')
        headers = list(next(sheet.values))
        def column(label):
            indices = [i for i, h in enumerate(headers) if str(h or '').strip().casefold() == label.casefold()]
            require(len(indices) == 1, 'Missing or duplicate column: ' + label)
            return indices[0]
        name_col, id_col = column('Full Name'), column('Student ID (College Reg. No.)')
        start_col, end_col = column('Start time'), column('Completion time')
        mapping = []
        for q in questions:
            target = normalized(q['text'])
            found = [i for i, h in enumerate(headers) if len(normalized(h)) >= 30 and
                     (normalized(h) == target or target.startswith(normalized(h)))]
            require(len(found) == 1, 'Question column cannot be matched uniquely: ' + q['text'][:90])
            mapping.append((q, found[0]))
        require(len({i for _, i in mapping}) == len(mapping), 'Question columns overlap. Check the worksheet headings.')
        records, errors, seen = [], [], set()
        for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(v is not None and v != '' for v in row):
                continue
            try:
                name = clean_text(str(row[name_col] or ''), 100, 'candidate name')
                sid = clean_text(str(row[id_col] or ''), 60, 'student ID')
                require(name and sid, 'Name and student ID are required.')
                key = student_key(sid)
                require(key not in seen, 'Duplicate student ID: ' + sid)
                seen.add(key)
                start, end = row[start_col], row[end_col]
                require(isinstance(start, datetime) and isinstance(end, datetime), 'Start/completion times must be Excel dates.')
                elapsed = (end - start).total_seconds() / 60
                require(0 <= elapsed <= 1440, 'Invalid completion time.')
                responses = {}
                for q, index in mapping:
                    value = row[index]
                    if value is None or str(value).strip() == '':
                        continue
                    matches = [i for i, option in enumerate(q['options']) if normalized(value) == normalized(option)]
                    require(len(matches) == 1, 'Answer does not match an option for ' + q['text'][:55] + ': ' + str(value)[:80])
                    responses[q['id']] = matches
                records.append({'name': name, 'studentId': sid, 'responses': responses,
                                'importedElapsedMinutes': round(elapsed, 4)})
            except ValueError as error:
                errors.append(f'Row {number}: {error}')
        require(records or errors, 'No candidate response rows found.')
        return records, errors, [{'column': headers[index], 'question': q['text']} for q, index in mapping]
    finally:
        book.close()


def preview_fingerprint(raw, questions, version):
    return hashlib.sha256(raw + json.dumps(questions, sort_keys=True).encode() + str(version).encode()).hexdigest()