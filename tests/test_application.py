import copy
import io
import json
import tempfile
import time
import unittest
import uuid
import zipfile
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from app import create_app
from domain import grade


def uid():
    return str(uuid.uuid4())


def question(kind='single', key=None):
    return {'id': 'q1', 'type': kind, 'text': 'Synthetic test question',
            'options': ['One', 'Two'] if kind != 'text' else [],
            'key': [1] if key is None and kind != 'text' else (key or []),
            'correctScore': 2, 'wrongScore': -1,
            'referenceAnswer': 'PRIVATE REFERENCE', 'reviewNote': 'PRIVATE REVIEW'}


class ApplicationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.password_hash = generate_password_hash('Test-only-password-123')

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent)
        self.url = 'sqlite:///' + str(Path(self.directory.name, 'test.db')).replace('\\', '/')
        self.settings = {'TESTING': True, 'APP_ENV': 'test', 'DATABASE_URL': self.url,
                         'SECRET_KEY': 'test-secret-not-production', 'ADMIN_PASSWORD_HASH': self.password_hash}
        self.app = create_app(self.settings)
        self.admin = self.app.test_client()
        csrf = self.admin.get('/api/auth/session').json['csrf']
        result = self.admin.post('/api/auth/login', json={'email': 'turnbrecruitment@gmail.com', 'password': 'Test-only-password-123'}, headers={'X-CSRF-Token': csrf})
        self.csrf = result.json['csrf']
        self.drive = uid()
        self.tests = {kind: uid() for kind in ('aptitude', 'language', 'prompt')}
        self.choice_orders = {}
        state = {'drives': [{'id': self.drive, 'name': 'Synthetic college 2026', 'finishedAt': None}],
                 'questionSets': {'aptitude': [question()], 'language': [question('text')], 'prompt': [question('text')]},
                 'sessions': [{'id': self.tests['aptitude'], 'driveId': self.drive, 'type': 'aptitude',
                               'questions': [question()], 'closedAt': None}]}
        response = self.put_state(state, 0)
        self.assertEqual(response.status_code, 200, response.json)
        if self._testMethodName != 'test_matching_key_is_not_exposed_and_choices_are_attempt_bound':
            candidate, candidate_csrf, _ = self.candidate(student='GD-SEED')
            submitted = self.answer(candidate, candidate_csrf, responses={}, submit=True)
            self.assertEqual(submitted.status_code, 200, submitted.json)
            shortlisted = self.admin.post(f'/api/admin/tests/{self.tests["aptitude"]}/shortlist', json={'value': True},
                                          headers={'X-CSRF-Token': self.csrf})
            self.assertEqual(shortlisted.status_code, 200, shortlisted.json)
            pools = self.admin.post(f'/api/admin/drives/{self.drive}/gd/pools', json={'poolSize': 7},
                                    headers={'X-CSRF-Token': self.csrf})
            self.assertEqual(pools.status_code, 200, pools.json)
        current = self.state()
        current['state']['sessions'].extend([
            {'id': self.tests[kind], 'driveId': self.drive, 'type': kind,
             'questions': [question('text')], 'closedAt': None}
            for kind in ('language', 'prompt')])
        response = self.put_state(current['state'], current['revision'])
        self.assertEqual(response.status_code, 200, response.json)

    def tearDown(self):
        self.app.extensions['store'].engine.dispose()
        self.directory.cleanup()

    def test_existing_drive_members_can_be_added_without_losing_evaluations(self):
        before = self.state()['state']['drives'][0]
        url = f'/api/admin/drives/{self.drive}/members'
        result = self.admin.post(url, json={'members': ['Alex', 'alex', 'Sam']},
                                 headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(result.status_code, 200, result.json)
        self.assertEqual(result.json['members'], ['Alex', 'Sam'])
        after = self.state()['state']['drives'][0]
        self.assertEqual(after.get('gd'), before.get('gd'))
        self.assertEqual(after.get('finalDecisions'), before.get('finalDecisions'))
        result = self.admin.post(url, json={'members': ['Taylor']},
                                 headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(result.json['members'], ['Alex', 'Sam', 'Taylor'])
        anonymous = self.app.test_client()
        anonymous_csrf = anonymous.get('/api/auth/session').json['csrf']
        self.assertEqual(anonymous.post(url, json={'members': ['Intruder']}, headers={'X-CSRF-Token': anonymous_csrf}).status_code, 401)

    def test_scheduled_links_hide_questions_and_block_early_starts_for_all_tests(self):
        current = self.state()
        drive_id = uid()
        available = int(time.time() * 1000) + 60000
        current['state']['drives'].append({'id': drive_id, 'name': 'Scheduled drive'})
        test_ids = {}
        for kind in ('aptitude', 'language', 'prompt'):
            test_ids[kind] = uid()
            current['state']['sessions'].append({'id': test_ids[kind], 'driveId': drive_id,
                'type': kind, 'questions': [question()], 'availableAt': available})
        saved = self.put_state(current['state'], current['revision'])
        self.assertEqual(saved.status_code, 200, saved.json)
        clients = {}
        for kind, tid in test_ids.items():
            client = self.app.test_client()
            csrf = client.get('/api/auth/session').json['csrf']
            clients[kind] = (client, csrf)
            with patch('app.now_ms', return_value=available - 1):
                meta = client.get(f'/api/candidate/{tid}').json['test']
                self.assertEqual(meta['availableAt'], available)
                self.assertEqual(meta['questions'], [])
                self.assertEqual(meta['attempts'], [])
                blocked = client.post(f'/api/candidate/{tid}/start', json={'name': 'Scheduled Candidate', 'studentId': 'SCHEDULE-1'},
                                      headers={'X-CSRF-Token': csrf})
                self.assertEqual(blocked.status_code, 400, blocked.json)
                self.assertIn('scheduled time', blocked.json['error'])
                if kind != 'aptitude':
                    identity = client.post(f'/api/candidate/{tid}/identity', json={'studentId': 'SCHEDULE-1'},
                                           headers={'X-CSRF-Token': csrf})
                    self.assertEqual(identity.status_code, 400)
        for kind, tid in test_ids.items():
            client, csrf = clients[kind]
            with patch('app.now_ms', return_value=available):
                result = client.post(f'/api/candidate/{tid}/start', json={'name': 'Scheduled Candidate', 'studentId': 'SCHEDULE-1'},
                                     headers={'X-CSRF-Token': csrf})
                self.assertEqual(result.status_code, 200, result.json)
                attempt = result.json['test']['attempts'][0]
                self.assertEqual(attempt['startedAt'], available)
                self.assertEqual(attempt['deadline'], available + result.json['test']['durationMinutes'] * 60000)
                self.assertTrue(result.json['test']['questions'])
        current = self.state()
        test = next(s for s in current['state']['sessions'] if s['id'] == test_ids['aptitude'])
        test['availableAt'] = available + 1000
        self.assertEqual(self.put_state(current['state'], current['revision']).status_code, 400)

    def put_state(self, state, revision):
        return self.admin.put('/api/admin/state', json={'state': state, 'revision': revision}, headers={'X-CSRF-Token': self.csrf})

    def candidate(self, kind='aptitude', name='Test Candidate', student='ID001'):
        if kind != 'aptitude':
            existing = self.state()['state']['sessions'][0]['attempts']
            if not any(a['studentId'] == student for a in existing):
                master, master_csrf, _ = self.candidate('aptitude', name, student)
                self.answer(master, master_csrf, submit=True)
        client = self.app.test_client()
        csrf = client.get('/api/auth/session').json['csrf']
        tid = self.tests[kind]
        response = client.post(f'/api/candidate/{tid}/start', json={'name': name, 'studentId': student}, headers={'X-CSRF-Token': csrf})
        self.assertEqual(response.status_code, 200, response.json)
        candidate_attempt = response.json['test']['attempts'][0]
        with self.app.extensions['store'].transaction() as tx:
            stored_attempt = next(a for a in tx.all_attempts(tid) if a['id'] == candidate_attempt['id'])
            self.choice_orders[id(client)] = copy.deepcopy(stored_attempt.get('optionOrder', {}))
        return client, csrf, candidate_attempt

    def answer(self, client, csrf, kind='aptitude', responses=None, sequence=1, submit=False):
        payload_responses = copy.deepcopy(responses if responses is not None else {'q1': [1]})
        if kind == 'aptitude':
            option_orders = self.choice_orders.get(id(client), {})
            for question_id, selected in list(payload_responses.items()):
                order = option_orders.get(question_id)
                if (order and isinstance(selected, list)
                        and all(isinstance(index, int) and not isinstance(index, bool) for index in selected)):
                    payload_responses[question_id] = [order.index(index) for index in selected]
        return client.post(f'/api/candidate/{self.tests[kind]}/answer', json={'responses': payload_responses, 'sequence': sequence, 'submit': submit}, headers={'X-CSRF-Token': csrf})

    def state(self):
        return self.admin.get('/api/admin/state').json

    def test_later_tests_resolve_aptitude_identity_and_reject_unknown_ids(self):
        guest = self.app.test_client()
        csrf = guest.get('/api/auth/session').json['csrf']
        language = self.tests['language']
        unknown = guest.post(f'/api/candidate/{language}/identity', json={'studentId': 'MISSING'}, headers={'X-CSRF-Token': csrf})
        self.assertEqual(unknown.status_code, 400)
        refused = guest.post(f'/api/candidate/{language}/start', json={'name': 'Other', 'studentId': 'MISSING'}, headers={'X-CSRF-Token': csrf})
        self.assertEqual(refused.status_code, 400)
        master, master_csrf, _ = self.candidate('aptitude', name='Canonical Name', student='MASTER-1')
        self.answer(master, master_csrf, submit=True)
        found = guest.post(f'/api/candidate/{language}/identity', json={'studentId': 'master-1'}, headers={'X-CSRF-Token': csrf})
        self.assertEqual(found.status_code, 200, found.json)
        self.assertEqual(found.json['name'], 'Canonical Name')
        started = guest.post(f'/api/candidate/{language}/start', json={'name': 'Spoofed Name', 'studentId': 'master-1'}, headers={'X-CSRF-Token': csrf})
        self.assertEqual(started.status_code, 200, started.json)
        attempt = started.json['test']['attempts'][0]
        self.assertEqual((attempt['name'], attempt['studentId']), ('Canonical Name', 'MASTER-1'))

    def test_language_shortlist_drives_new_gd_pools_and_clear_removes_scores(self):
        master, master_csrf, _ = self.candidate('aptitude', name='Language Candidate', student='LANG-GD')
        self.answer(master, master_csrf, submit=True)
        shortlisted = self.admin.post(f'/api/admin/tests/{self.tests["aptitude"]}/shortlist', json={'value': True}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(shortlisted.status_code, 200)
        language, language_csrf, attempt = self.candidate('language', name='Language Candidate', student='LANG-GD')
        self.answer(language, language_csrf, 'language', responses={'q1': 'Essay'}, submit=True)
        mark = self.admin.post(f'/api/admin/tests/{self.tests["language"]}/marks', json={'attemptId': attempt['id'], 'questionId': 'q1', 'score': 2}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(mark.status_code, 200, mark.json)
        short = self.admin.post(f'/api/admin/tests/{self.tests["language"]}/shortlist', json={'value': True}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(short.status_code, 200, short.json)
        pools = self.admin.post(f'/api/admin/drives/{self.drive}/gd/pools', json={'poolSize': 7}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(pools.status_code, 200, pools.json)
        self.assertEqual(pools.json['gd']['source'], 'language')
        self.assertIn('lang-gd', pools.json['gd']['candidates'])
        evaluator = self.admin.post(f'/api/admin/drives/{self.drive}/gd/evaluator', json={'poolIndex': 0, 'evaluatorName': 'Test Evaluator'}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(evaluator.status_code, 200, evaluator.json)
        score = self.admin.post(f'/api/admin/drives/{self.drive}/gd/score', json={'studentKey': 'lang-gd', 'field': 'initiative', 'score': 4}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(score.status_code, 200, score.json)
        cleared = self.admin.post(f'/api/admin/drives/{self.drive}/gd/clear', json={'confirm': True}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(cleared.status_code, 200, cleared.json)
        drive = next(d for d in self.state()['state']['drives'] if d['id'] == self.drive)
        self.assertNotIn('gd', drive)

    def test_security_and_no_question_leakage(self):
        guest = self.app.test_client()
        self.assertEqual(guest.get('/api/admin/state').status_code, 401)
        self.assertEqual(guest.get(f'/api/admin/tests/{self.tests["aptitude"]}/excel').status_code, 401)
        self.assertEqual(guest.post('/api/auth/logout', json={}).status_code, 403)
        public = guest.get('/api/candidate/' + self.tests['aptitude']).json
        self.assertEqual(public['test']['questions'], [])
        client, _, _ = self.candidate(student='SECURITY')
        public = client.get('/api/candidate/' + self.tests['aptitude']).json
        serialized = json.dumps(public)
        self.assertNotIn('PRIVATE REFERENCE', serialized)
        self.assertNotIn('PRIVATE REVIEW', serialized)
        self.assertNotIn('"key"', serialized)
        self.assertNotIn('token_hash', serialized)
        self.assertEqual(guest.get('/healthz').headers['X-Frame-Options'], 'DENY')

    def test_individual_durations_and_duplicate_student(self):
        for kind, minutes in [('aptitude', 15), ('language', 20), ('prompt', 30)]:
            _, _, a = self.candidate(kind, student='DURATION-' + kind)
            self.assertEqual(a['deadline'] - a['startedAt'], minutes * 60000)
        other = self.app.test_client()
        csrf = other.get('/api/auth/session').json['csrf']
        response = other.post('/api/candidate/' + self.tests['aptitude'] + '/start', json={'name': 'Other', 'studentId': 'duration-aptitude'}, headers={'X-CSRF-Token': csrf})
        self.assertEqual(response.status_code, 400)

    def test_submission_idempotency_and_save_sequence(self):
        c, csrf, _ = self.candidate(student='SEQUENCE')
        first_save = self.answer(c, csrf, sequence=2)
        self.assertEqual(first_save.status_code, 200)
        expected_displayed_answer = first_save.json['test']['attempts'][0]['responses']['q1']
        a = self.answer(c, csrf, responses={'q1': [0]}, sequence=1).json['test']['attempts'][0]
        self.assertEqual(a['responses']['q1'], expected_displayed_answer)
        result = self.answer(c, csrf, sequence=3, submit=True)
        original = result.json['test']['attempts'][0]
        later = self.answer(c, csrf, responses={'q1': [0]}, sequence=4, submit=True).json['test']['attempts'][0]
        self.assertEqual(original['submittedAt'], later['submittedAt'])
        self.assertEqual(later['responses']['q1'], expected_displayed_answer)

    def test_matching_key_is_not_exposed_and_choices_are_attempt_bound(self):
        state = self.state()
        test = next(s for s in state['state']['sessions'] if s['type'] == 'aptitude')
        test['questions'] = [{'id': 'match1', 'type': 'matching', 'text': 'Match synthetic items',
                             'pairs': [{'left': 'A', 'right': 'a'}, {'left': 'B', 'right': 'b'}],
                             'correctScore': 2, 'wrongScore': 0}]
        self.assertEqual(self.put_state(state['state'], state['revision']).status_code, 200)
        client, csrf, _ = self.candidate(student='MATCHING')
        public = client.get('/api/candidate/' + self.tests['aptitude']).json['test']['questions'][0]
        self.assertNotIn('pairs', public)
        self.assertEqual(public['leftItems'], ['A', 'B'])
        tokens = {p['label']: p['id'] for p in public['rightChoices']}
        result = self.answer(client, csrf, responses={'match1': [tokens['a'], tokens['b']]}, submit=True)
        self.assertEqual(result.status_code, 200, result.json)
        with self.app.extensions['store'].transaction() as tx:
            a = tx.all_attempts(self.tests['aptitude'])[0]
            self.assertEqual(grade(tx.test(self.tests['aptitude'])['questions'], a)['total'], 2)

    def test_timeout_rejects_late_payload_and_resume_survives_restart(self):
        c, csrf, a = self.candidate(student='TIMEOUT')
        saved = self.answer(c, csrf)
        saved_displayed_answer = saved.json['test']['attempts'][0]['responses']['q1']
        with self.app.extensions['store'].transaction() as tx:
            next(x for x in tx.all_attempts(self.tests['aptitude']) if x['id'] == a['id'])['deadline'] = int(time.time() * 1000) - 100
        result = self.answer(c, csrf, responses={'q1': [0]}, sequence=3, submit=True)
        locked = result.json['test']['attempts'][0]
        self.assertEqual(locked['reason'], 'timeout')
        self.assertEqual(locked['responses']['q1'], saved_displayed_answer)
        restarted = create_app(self.settings)
        with restarted.extensions['store'].transaction() as tx:
            self.assertEqual(next(x for x in tx.all_attempts(self.tests['aptitude']) if x['id'] == a['id'])['responses']['q1'], [1])
        restarted.extensions['store'].engine.dispose()

    def test_admin_stop_saves_blank_and_partial_answers(self):
        c, csrf, _ = self.candidate(student='PARTIAL')
        self.answer(c, csrf)
        self.candidate(student='BLANK')
        state = self.state()
        next(s for s in state['state']['sessions'] if s['type'] == 'aptitude')['closedAt'] = 1
        result = self.put_state(state['state'], state['revision'])
        self.assertEqual(result.status_code, 200, result.json)
        test = next(s for s in result.json['state']['sessions'] if s['type'] == 'aptitude')
        stopped = [a for a in test['attempts'] if a['studentId'] in {'PARTIAL', 'BLANK'}]
        self.assertTrue(all(a['reason'] == 'admin' for a in stopped))
        self.assertEqual(len(next(a for a in stopped if a['studentId'] == 'BLANK')['responses']), 0)
        self.assertEqual(c.get('/api/candidate/' + self.tests['aptitude']).json['test']['attempts'][0]['reason'], 'admin')

    def test_question_lock_and_admin_revision_conflict(self):
        state = self.state()
        self.candidate(student='LOCKED')
        next(s for s in state['state']['sessions'] if s['type'] == 'aptitude')['questions'][0]['text'] = 'Changed after start'
        self.assertEqual(self.put_state(state['state'], state['revision']).status_code, 400)
        state = self.state()
        state['state']['drives'][0]['name'] = 'Renamed'
        self.assertEqual(self.put_state(state['state'], state['revision']).status_code, 200)
        self.assertEqual(self.put_state(state['state'], state['revision']).status_code, 409)

    def test_marks_filters_bulk_shortlist_and_excel(self):
        c, csrf, a = self.candidate(name='=HYPERLINK("https://invalid")', student='00001')
        self.answer(c, csrf, submit=True)
        result = self.admin.post(f'/api/admin/tests/{self.tests["aptitude"]}/shortlist', json={'minMarks': 2, 'minTime': 0}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(result.json['changed'], 1)
        exported = self.admin.get(f'/api/admin/tests/{self.tests["aptitude"]}/excel?minMarks=2&minTime=0')
        book = load_workbook(io.BytesIO(exported.data))
        self.assertEqual(book.active.cell(2, 3).value, '00001')
        self.assertEqual(book.active.cell(2, 4).value, 2)
        self.assertEqual(book.active.cell(2, 2).data_type, 's')
        self.assertEqual(book.active.cell(2, 7).value, 'Yes')
        reverted = self.admin.post(f'/api/admin/tests/{self.tests["aptitude"]}/shortlist', json={'value': False}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(reverted.json['changed'], 2)
        cleared = load_workbook(io.BytesIO(self.admin.get(f'/api/admin/tests/{self.tests["aptitude"]}/excel').data))
        self.assertEqual(cleared.active.cell(2, 7).value, 'No')
        self.assertEqual(self.admin.get(f'/api/admin/tests/{self.tests["aptitude"]}/excel?minTime=-1').status_code, 400)

    def test_manual_review_blocks_shortlist_then_unlocks(self):
        c, csrf, a = self.candidate('language')
        self.answer(c, csrf, 'language', responses={'q1': 'Essay'}, submit=True)
        url = f'/api/admin/tests/{self.tests["language"]}'
        response = self.admin.post(url + '/shortlist', json={}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(response.json['skipped'], 1)
        marked = self.admin.post(url + '/marks', json={'attemptId': a['id'], 'questionId': 'q1', 'score': 1.5}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(marked.json['grade']['total'], 1.5)
        response = self.admin.post(url + '/shortlist', json={}, headers={'X-CSRF-Token': self.csrf})
        self.assertEqual(response.json['changed'], 1)

    def test_prompt_submissions_leave_session_open_and_exports_available(self):
        first, csrf1, a = self.candidate('prompt', name='First Candidate', student='P1')
        second, csrf2, _ = self.candidate('prompt', name='Second Candidate', student='P2')
        url = f'/api/admin/tests/{self.tests["prompt"]}/word'
        self.assertEqual(self.admin.get(url).status_code, 400)
        self.answer(first, csrf1, 'prompt', {'q1': 'Long response\nwith <markup> & symbols'}, submit=True)
        self.assertEqual(self.admin.get(url).status_code, 400)
        self.answer(second, csrf2, 'prompt', {'q1': ''}, submit=True)
        response = self.admin.get(url)
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            self.assertEqual(set(archive.namelist()), {'First Candidate_P1.docx', 'Second Candidate_P2.docx'})
            doc = Document(io.BytesIO(archive.read('First Candidate_P1.docx')))
            self.assertIn('with <markup> & symbols', '\n'.join(p.text for p in doc.paragraphs))
            doc = Document(io.BytesIO(archive.read('Second Candidate_P2.docx')))
            self.assertIn('(No answer entered)', '\n'.join(p.text for p in doc.paragraphs))

    def test_prompt_late_join_timeout_and_admin_stop_are_independent(self):
        first, csrf1, a = self.candidate('prompt', student='FIRST')
        self.answer(first, csrf1, 'prompt', {'q1': 'First response'}, submit=True)
        second, csrf2, b = self.candidate('prompt', student='LATER')
        with self.app.extensions['store'].transaction() as tx:
            attempt = next(item for item in tx.all_attempts(self.tests['prompt']) if item['id'] == b['id'])
            attempt['deadline'] = int(time.time() * 1000) - 1000
        current = self.state()
        test = next(t for t in current['state']['sessions'] if t['id'] == self.tests['prompt'])
        self.assertIsNone(test['closedAt'])
        expired = next(item for item in test['attempts'] if item['id'] == b['id'])
        self.assertEqual(expired['reason'], 'timeout')
        self.assertTrue(expired['submittedAt'])
        third, csrf3, c = self.candidate('prompt', student='AFTER-TIMEOUT')
        current = self.state()
        test = next(t for t in current['state']['sessions'] if t['id'] == self.tests['prompt'])
        test['closedAt'] = int(time.time() * 1000)
        stopped = self.put_state(current['state'], current['revision'])
        self.assertEqual(stopped.status_code, 200, stopped.json)
        test = next(t for t in stopped.json['state']['sessions'] if t['id'] == self.tests['prompt'])
        self.assertTrue(test['closedAt'])
        active = next(item for item in test['attempts'] if item['id'] == c['id'])
        self.assertEqual(active['reason'], 'admin')
        self.assertTrue(active['submittedAt'])
        blocked = self.app.test_client()
        csrf = blocked.get('/api/auth/session').json['csrf']
        result = blocked.post(f'/api/candidate/{self.tests["prompt"]}/start', json={'studentId': 'FIRST'}, headers={'X-CSRF-Token': csrf})
        self.assertEqual(result.status_code, 400)

    def test_fifty_concurrent_starts_saves_and_submissions(self):
        def worker(number):
            client = self.app.test_client()
            csrf = client.get('/api/auth/session').json['csrf']
            headers = {'X-CSRF-Token': csrf}
            start = client.post(f'/api/candidate/{self.tests["aptitude"]}/start', json={'name': f'Synthetic {number}', 'studentId': f'LOAD{number:03}'}, headers=headers)
            if start.status_code != 200:
                return start.status_code
            saved = self.answer(client, csrf, sequence=1)
            submitted = self.answer(client, csrf, sequence=2, submit=True)
            return submitted.status_code if saved.status_code == 200 else saved.status_code
        begin = time.monotonic()
        with ThreadPoolExecutor(max_workers=50) as pool:
            results = list(pool.map(worker, range(50)))
        self.assertEqual(results, [200] * 50)
        state = self.state()
        test = next(s for s in state['state']['sessions'] if s['type'] == 'aptitude')
        self.assertEqual(sum(a['name'].startswith('Synthetic ') for a in test['attempts']), 50)
        self.assertTrue(all(a['submittedAt'] for a in test['attempts']))
        print(f'50-client durable-store concurrency test completed in {time.monotonic()-begin:.2f}s')


    def test_manual_aptitude_shortlist_requires_remarks_and_is_atomic(self):
        client, csrf, attempt = self.candidate(student='MANUAL')
        self.answer(client, csrf, responses={'q1': [0]}, submit=True)
        url = f'/api/admin/tests/{self.tests["aptitude"]}/manual-shortlist'
        headers = {'X-CSRF-Token': self.csrf}
        change = {'attemptId': attempt['id'], 'shortlisted': True, 'remarks': ''}
        self.assertEqual(self.admin.post(url, json={'changes': [change]}, headers=headers).status_code, 400)
        change['remarks'] = 'Solved the difficult question with a strong approach.'
        invalid = {'attemptId': 'missing', 'shortlisted': True, 'remarks': 'Invalid candidate'}
        self.assertEqual(self.admin.post(url, json={'changes': [change, invalid]}, headers=headers).status_code, 400)
        stored = next(a for a in self.state()['state']['sessions'][0]['attempts'] if a['id'] == attempt['id'])
        self.assertFalse(stored['shortlisted'])
        self.assertEqual(self.admin.post(url, json={'changes': [change]}, headers=headers).status_code, 200)
        stored = next(a for a in self.state()['state']['sessions'][0]['attempts'] if a['id'] == attempt['id'])
        self.assertTrue(stored['shortlisted'])
        self.assertTrue(stored['manualShortlisted'])
        self.assertEqual(stored['shortlistRemarks'], change['remarks'])
        public = client.get('/api/candidate/' + self.tests['aptitude']).json
        self.assertNotIn('shortlistRemarks', json.dumps(public))
        change.update(shortlisted=False, remarks='')
        self.assertEqual(self.admin.post(url, json={'changes': [change]}, headers=headers).status_code, 200)
        stored = next(a for a in self.state()['state']['sessions'][0]['attempts'] if a['id'] == attempt['id'])
        self.assertFalse(stored['shortlisted'])
        self.assertEqual(stored['shortlistRemarks'], '')

    def test_gd_random_allocation_and_both_workbook_exports(self):
        from unittest.mock import patch
        headers = {'X-CSRF-Token': self.csrf}
        for number in range(6):
            client, csrf, _ = self.candidate(name=f'Candidate {number}', student=f'POOL-{number}')
            self.answer(client, csrf, submit=True)
        self.admin.post(f'/api/admin/tests/{self.tests["aptitude"]}/shortlist', json={}, headers=headers)
        with self.app.extensions['store'].transaction() as tx:
            expected = [a['studentId'].casefold() for a in tx.all_attempts(self.tests['aptitude']) if a.get('shortlisted')]
        with patch('app.secrets.SystemRandom') as rng:
            rng.return_value.shuffle.side_effect = lambda roster: roster.reverse()
            result = self.admin.post(f'/api/admin/drives/{self.drive}/gd/pools', json={'poolSize': 3}, headers=headers)
            rng.return_value.shuffle.assert_called_once()
        self.assertEqual(result.status_code, 200, result.json)
        gd = result.json['gd']
        self.assertEqual([key for pool in gd['pools'] for key in pool['members']], list(reversed(expected)))
        self.assertTrue(all(len(pool['members']) <= 10 for pool in gd['pools']))
        url = f'/api/admin/drives/{self.drive}/gd'
        self.admin.post(url + '/evaluator', json={'poolIndex': 0, 'evaluatorName': 'Evaluator A'}, headers=headers)
        key = gd['pools'][0]['members'][0]
        self.admin.post(url + '/score', json={'studentKey': key, 'field': 'initiative', 'score': 4}, headers=headers)
        roster = load_workbook(io.BytesIO(self.admin.get(url + '/export?mode=roster').data))
        self.assertEqual(roster.sheetnames, ['GD pools'])
        self.assertEqual(roster.active.max_row, len(expected) + 1)
        scores = load_workbook(io.BytesIO(self.admin.get(url + '/export?mode=scores').data))
        self.assertEqual(scores.sheetnames, [pool['name'] for pool in gd['pools']])
        self.assertEqual(scores.worksheets[0].cell(2, 3).value, 'Evaluator A')
        self.assertEqual(scores.worksheets[0].cell(2, 7).value, 4)
        self.assertEqual(scores.worksheets[0].cell(2, 11).value, 4)
        self.assertEqual(self.app.test_client().get(url + '/export').status_code, 401)

    def test_manual_interview_decisions_and_exports_include_evaluator_and_interviewer(self):
        headers = {'X-CSRF-Token': self.csrf}
        student = 'GD-SEED'
        base = f'/api/admin/drives/{self.drive}'
        self.admin.post(base + '/gd/evaluator', json={'poolIndex': 0, 'evaluatorName': 'GD Evaluator'}, headers=headers)
        url = base + '/candidate-decision'
        data = {'studentId': student, 'interviewShortlisted': True}
        self.assertEqual(self.admin.post(url, json=data, headers=headers).status_code, 400)
        data.update(interviewRemarks='Strong reasoning in the group discussion.', interviewer='Interview Evaluator')
        response = self.admin.post(url, json=data, headers=headers)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertTrue(response.json['decision']['interviewShortlisted'])
        response = self.admin.post(url, json={'studentId': student, 'interviewRemarks': ''}, headers=headers)
        self.assertEqual(response.status_code, 400)
        exported = self.admin.post(base + '/final/export', json={'studentIds': [student]}, headers=headers)
        sheet = load_workbook(io.BytesIO(exported.data)).active
        values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
        self.assertEqual(values['GD evaluator'], 'GD Evaluator')
        self.assertEqual(values['Interview shortlisted'], 'Yes')
        self.assertEqual(values['Interview remarks'], data['interviewRemarks'])
        self.assertEqual(values['Interviewer'], data['interviewer'])
        self.admin.post(base + '/final-shortlist', json={'studentIds': [student], 'value': True}, headers=headers)
        decision = self.state()['state']['drives'][0]['finalDecisions']['gd-seed']
        self.assertEqual(decision['interviewRemarks'], '')
        self.assertFalse(decision['interviewManual'])

    def test_response_comparisons_apply_to_shortlist_and_excel(self):
        from urllib.parse import urlencode
        headers = {'X-CSRF-Token': self.csrf}
        for index, score in enumerate((10, 20, 30)):
            client, csrf, attempt = self.candidate(student=f'FILTER-{index}')
            self.answer(client, csrf, submit=True)
            with self.app.extensions['store'].transaction() as tx:
                stored = next(a for a in tx.all_attempts(self.tests['aptitude']) if a['id'] == attempt['id'])
                stored['importedScore'] = score
                stored['importedElapsedMinutes'] = index + 1
        filters = {'marksOp': 'between', 'minMarks': 10, 'maxMarks': 20,
                   'timeOp': 'lte', 'minTime': 2}
        url = f'/api/admin/tests/{self.tests["aptitude"]}'
        result = self.admin.post(url + '/shortlist', json={**filters, 'value': True}, headers=headers)
        self.assertEqual(result.status_code, 200, result.json)
        attempts = self.state()['state']['sessions'][0]['attempts']
        self.assertEqual({a['studentId'] for a in attempts if a['studentId'].startswith('FILTER-') and a.get('shortlisted')},
                         {'FILTER-0', 'FILTER-1'})
        exported = self.admin.get(url + '/excel?' + urlencode(filters))
        self.assertEqual(exported.status_code, 200)
        sheet = load_workbook(io.BytesIO(exported.data)).active
        self.assertEqual({sheet.cell(i, 3).value for i in range(2, sheet.max_row + 1)}, {'FILTER-0', 'FILTER-1'})
        for op, expected in [('gt', ['FILTER-2']), ('lt', ['FILTER-0']), ('eq', ['FILTER-1'])]:
            params = {'marksOp': op, 'minMarks': 20}
            book = load_workbook(io.BytesIO(self.admin.get(url + '/excel?' + urlencode(params)).data))
            actual = [row[2] for row in book.active.iter_rows(min_row=2, values_only=True) if str(row[2]).startswith('FILTER-')]
            self.assertEqual(actual, expected)
        bad = self.admin.get(url + '/excel?' + urlencode({'marksOp': 'between', 'minMarks': 20, 'maxMarks': 10}))
        self.assertEqual(bad.status_code, 400)

    def test_performance_report_uses_original_answer_indices_after_shuffle(self):
        for kind in ('aptitude', 'language'):
            if kind == 'language':
                current = self.state()
                test = next(t for t in current['state']['sessions'] if t['type'] == kind)
                test['questions'] = [question()]
                self.assertEqual(self.put_state(current['state'], current['revision']).status_code, 200)
            for display_index, expected_score, expected_status in ((0, 2, 'Correct'), (1, -1, 'Wrong')):
                student = f'REPORT-{kind}-{display_index}'
                client, csrf, attempt = self.candidate(kind, student=student)
                with self.app.extensions['store'].transaction() as tx:
                    stored = next(a for a in tx.all_attempts(self.tests[kind]) if a['id'] == attempt['id'])
                    stored['optionOrder']['q1'] = [1, 0]
                result = client.post(f'/api/candidate/{self.tests[kind]}/answer',
                                    json={'responses': {'q1': [display_index]}, 'sequence': 1, 'submit': True},
                                    headers={'X-CSRF-Token': csrf})
                self.assertEqual(result.status_code, 200, result.json)
                report = self.admin.get(f'/api/admin/drives/{self.drive}/candidate-responses?studentId={student}')
                self.assertEqual(report.status_code, 200, report.json)
                section = next(s for s in report.json['sections'] if s['type'] == kind)
                item = section['questions'][0]
                selected = [c for c in item['choices'] if c['selected']]
                self.assertEqual([c['text'] for c in selected], ['Two' if display_index == 0 else 'One'])
                self.assertEqual(selected[0]['correct'], display_index == 0)
                self.assertEqual(item['score'], expected_score)
                self.assertEqual(item['status'], expected_status)
                self.assertEqual(section['score'], expected_score)

    def test_performance_report_prompt_and_split_final_scores(self):
        headers = {'X-CSRF-Token': self.csrf}
        client, csrf, _ = self.candidate('prompt', student='GD-SEED')
        self.answer(client, csrf, kind='prompt', responses={'q1': 'Role: analyst. Explain the result clearly.\nUse examples.'}, submit=True)
        report = self.admin.get(f'/api/admin/drives/{self.drive}/candidate-responses?studentId=GD-SEED')
        self.assertEqual(report.status_code, 200, report.json)
        prompt = next(s for s in report.json['sections'] if s['type'] == 'prompt')
        self.assertIn('Role: analyst.', prompt['response'])
        self.assertIn('Use examples.', prompt['response'])
        current = self.state()
        language = next(s for s in current['state']['sessions'] if s['type'] == 'language')
        language['questions'] = [{**question(), 'id': f'q{i}', 'correctScore': 1, 'wrongScore': 0} for i in range(25)]
        self.assertEqual(self.put_state(current['state'], current['revision']).status_code, 200)
        language_client, language_csrf, attempt = self.candidate('language', student='GD-SEED')
        with self.app.extensions['store'].transaction() as tx:
            stored = next(a for a in tx.all_attempts(self.tests['language']) if a['id'] == attempt['id'])
            stored['responses'] = {f'q{i}': [1] for i in range(25)}
            stored['submittedAt'] = int(time.time() * 1000)
        self.admin.post(f'/api/admin/drives/{self.drive}/final-shortlist', json={'studentIds': ['GD-SEED'], 'value': True}, headers=headers)
        exported = self.admin.post(f'/api/admin/drives/{self.drive}/final/export',
                                   json={'studentIds': ['GD-SEED'], 'mode': 'interview'}, headers=headers)
        self.assertEqual(exported.status_code, 200)
        sheet = load_workbook(io.BytesIO(exported.data)).active
        values = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
        self.assertEqual(values['Language score (20)'], 20)
        self.assertEqual(values['Reasoning score (5)'], 5)
        self.assertEqual(values['Interview shortlisted'], 'Yes')
        self.admin.post(f'/api/admin/drives/{self.drive}/final-shortlist', json={'studentIds': ['GD-SEED'], 'value': False}, headers=headers)
        empty = self.admin.post(f'/api/admin/drives/{self.drive}/final/export',
                               json={'studentIds': ['GD-SEED'], 'mode': 'interview'}, headers=headers)
        self.assertEqual(empty.status_code, 400)

    def test_filled_response_import_preview_validation_and_preservation(self):
        from datetime import datetime, timedelta
        from openpyxl import Workbook
        current = self.state()
        drive = uid()
        current['state']['drives'].append({'id': drive, 'name': 'Import answers'})
        self.assertEqual(self.put_state(current['state'], current['revision']).status_code, 200)
        url = f'/api/admin/drives/{drive}/aptitude/responses-import'
        # Long headings are truncated in the actual Microsoft Forms export.
        text = 'A synthetic multiple choice question with a long heading for import matching'
        current = self.state()
        current['state']['questionSets']['aptitude'] = [{**question(), 'text': text}]
        self.assertEqual(self.put_state(current['state'], current['revision']).status_code, 200)
        def workbook(answer='Two', duplicate=False):
            book=Workbook();sheet=book.active;sheet.title='Input Data'
            sheet.append(['Full Name','Student ID (College Reg. No.)','Start time','Completion time',text[:60]])
            start=datetime(2026,9,25,12,0)
            sheet.append(['Imported Student','IMP1',start,start+timedelta(minutes=7),answer])
            if duplicate: sheet.append(['Duplicate','imp1',start,start+timedelta(minutes=7),'One'])
            stream=io.BytesIO();book.save(stream);return stream.getvalue()
        def post(raw, **fields):
            return self.admin.post(url,data={'file':(io.BytesIO(raw),'answers.xlsx'),**fields},headers={'X-CSRF-Token':self.csrf})
        raw=workbook()
        preview=post(raw);self.assertEqual(preview.status_code,200,preview.json)
        self.assertEqual(preview.json['records'][0]['score'],2)
        self.assertEqual(preview.json['errors'],[])
        self.assertFalse(any(t['driveId']==drive for t in self.state()['state']['sessions']))
        saved=post(raw,confirm='true',fingerprint=preview.json['fingerprint'])
        self.assertEqual(saved.status_code,200,saved.json)
        self.assertEqual(post(raw,confirm='true',fingerprint=preview.json['fingerprint']).status_code,400)
        current=self.state();test=next(t for t in current['state']['sessions'] if t['driveId']==drive)
        self.assertEqual(len(test['attempts']),1)
        self.assertEqual(test['attempts'][0]['responses'],{'q1':[1]})
        self.assertFalse(test['attempts'][0]['shortlisted'])
        with self.app.extensions['store'].transaction() as tx:
            a=tx.all_attempts(test['id'])[0];a['shortlisted']=True;a['shortlistRemarks']='Keep existing decision'
            d=next(d for d in tx.state['drives'] if d['id']==drive);d['finalDecisions']={'imp1':{'outcome':'selected'}}
        raw=workbook('One');preview=post(raw)
        self.assertEqual(preview.json['records'][0]['score'],-1)
        self.assertEqual(preview.json['records'][0]['action'],'Update')
        saved=post(raw,confirm='true',fingerprint=preview.json['fingerprint']);self.assertEqual(saved.status_code,200,saved.json)
        current=self.state();test=next(t for t in current['state']['sessions'] if t['driveId']==drive)
        self.assertEqual(len(test['attempts']),1);self.assertTrue(test['attempts'][0]['shortlisted'])
        self.assertEqual(test['attempts'][0]['shortlistRemarks'],'Keep existing decision')
        self.assertEqual(next(d for d in current['state']['drives'] if d['id']==drive)['finalDecisions']['imp1']['outcome'],'selected')
        for invalid in (workbook('Not an option'),workbook(duplicate=True)):
            preview=post(invalid);self.assertTrue(preview.json['errors'])
            self.assertEqual(post(invalid,confirm='true',fingerprint=preview.json['fingerprint']).status_code,400)
        guest=self.app.test_client();guest_csrf=guest.get('/api/auth/session').json['csrf']
        self.assertEqual(guest.post(url,json={},headers={'X-CSRF-Token':guest_csrf}).status_code,401)
        online=f'/api/admin/drives/{self.drive}/aptitude/responses-import'
        result=self.admin.post(online,data={'file':(io.BytesIO(workbook()),'answers.xlsx')},headers={'X-CSRF-Token':self.csrf})
        self.assertEqual(result.status_code,400)

    def test_response_column_filters_share_shortlist_and_export_behavior(self):
        from urllib.parse import urlencode
        headers = {'X-CSRF-Token': self.csrf}
        first, csrf1, first_attempt = self.candidate(name='Header Alice', student='HEADER-A')
        second, csrf2, second_attempt = self.candidate(name='Header Bob', student='HEADER-B')
        self.answer(first, csrf1, submit=True)
        self.answer(second, csrf2, submit=True)
        with self.app.extensions['store'].transaction() as tx:
            for attempt in tx.all_attempts(self.tests['aptitude']):
                if attempt['id'] in {first_attempt['id'], second_attempt['id']}:
                    attempt['importedScore'] = 20
                    attempt['importedElapsedMinutes'] = 3
        filters = {'driveName': {'text': 'Synthetic'}, 'name': {'text': 'alice'},
                   'studentId': {'text': 'HEADER'}, 'marks': {'op': 'gte', 'value': 15},
                   'minutes': {'op': 'between', 'value': 2, 'upper': 4}, 'shortlisted': {'values': ['No']}}
        url = f'/api/admin/tests/{self.tests["aptitude"]}'
        result = self.admin.post(url + '/shortlist', json={'columnFilters': filters, 'value': True}, headers=headers)
        self.assertEqual(result.status_code, 200, result.json)
        self.assertEqual(result.json['changed'], 1)
        filters['shortlisted']['values'] = ['Yes']
        exported = self.admin.get(url + '/excel?' + urlencode({'columnFilters': json.dumps(filters)}))
        self.assertEqual(exported.status_code, 200)
        sheet = load_workbook(io.BytesIO(exported.data)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 3).value, 'HEADER-A')
        invalid = self.admin.post(url + '/shortlist', json={'columnFilters': {'unknown': {'text': 'x'}}}, headers=headers)
        self.assertEqual(invalid.status_code, 400)
        from exports import results_rows
        questions = [{**question(), 'id': f'section-{i}', 'correctScore': 1, 'wrongScore': 0} for i in range(25)]
        language = {'type': 'language', 'questions': questions}
        candidates = [{'name': 'Alice', 'studentId': 'A', 'startedAt': 1, 'submittedAt': 60001,
                       'responses': {q['id']: [1] for q in questions}, 'shortlisted': True}]
        parts_filter = {'columnFilters': {'languageScore': {'op': 'eq', 'value': 20},
                                         'reasoningScore': {'op': 'gte', 'value': 5}}}
        self.assertEqual(len(results_rows(language, candidates, filters=parts_filter)), 1)
        parts_filter['columnFilters']['reasoningScore']['value'] = 6
        self.assertEqual(results_rows(language, candidates, filters=parts_filter), [])

    def test_prompt_workbook_has_unique_candidate_sheets_and_preserves_long_responses(self):
        name = 'Candidate / With * Long Name For Export'
        first, csrf1, _ = self.candidate('prompt', name=name, student='BOOK-1')
        second, csrf2, _ = self.candidate('prompt', name=name, student='BOOK-2')
        url = f'/api/admin/tests/{self.tests["prompt"]}/prompt-excel'
        self.assertEqual(self.app.test_client().get(url).status_code, 401)
        self.assertEqual(self.admin.get(url).status_code, 400)
        response_text = '=literal prompt\n' + 'Detailed response. ' * 3000
        self.answer(first, csrf1, 'prompt', {'q1': response_text}, submit=True)
        self.assertEqual(self.admin.get(url).status_code, 400)
        self.answer(second, csrf2, 'prompt', {'q1': ''}, submit=True)
        response = self.admin.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn('_Prompt%20Responses.xlsx', response.headers['Content-Disposition'].replace(' ', '%20'))
        book = load_workbook(io.BytesIO(response.data))
        self.assertEqual(len(book.sheetnames), 2)
        self.assertEqual(len(set(n.casefold() for n in book.sheetnames)), 2)
        self.assertTrue(all(len(n) <= 31 and not any(c in n for c in '[]:*?/\\') for n in book.sheetnames))
        first_sheet = next(s for s in book if s.cell(4, 2).value == 'BOOK-1')
        actual = ''.join(row[1].value for row in first_sheet if str(row[0].value).startswith('Response'))
        self.assertEqual(actual, response_text.strip())
        self.assertTrue(all(row[1].data_type == 's' for row in first_sheet if str(row[0].value).startswith('Response')))
        blank_sheet = next(s for s in book if s.cell(4, 2).value == 'BOOK-2')
        self.assertIn('(No answer entered)', [row[1].value for row in blank_sheet])

    def test_language_component_scores_use_original_question_ids(self):
        from exports import language_parts
        questions = [{'id': f'q{i}'} for i in range(25)]
        grading = {'results': [{'id': q['id'], 'score': 1 if i < 20 else 0.5} for i, q in enumerate(questions)]}
        grading['results'].reverse()
        self.assertEqual(language_parts({'questions': questions}, grading), [20, 2.5])
        grading['results'][0]['score'] = None
        self.assertEqual(language_parts({'questions': questions}, grading), [20, None])


if __name__ == '__main__':
    unittest.main()

