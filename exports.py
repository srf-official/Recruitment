"""Generate real XLSX and individual DOCX documents on the authenticated server."""
import io
import re
import zipfile

from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from domain import grade, require


def filename(value):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value)).rstrip('. ')[:160] or 'candidate'
    return '_' + name if re.match(r'^(con|prn|aux|nul|com[1-9]|lpt[1-9])(?:\.|$)', name, re.I) else name


def minutes(a):
    if isinstance(a.get('importedElapsedMinutes'), (int, float)):
        return round(max(0, a['importedElapsedMinutes']), 4)
    return round(max(0, (a['submittedAt'] - a['startedAt']) / 60000), 4) if a.get('submittedAt') else None


def number_filter(raw, operator='gte', upper=None):
    require(operator in {'gte', 'gt', 'lte', 'lt', 'eq', 'between', 'blank', 'notblank'}, 'Invalid numeric comparison.')
    if operator in {'blank', 'notblank'}:
        return lambda value: (value is None) if operator == 'blank' else (value is not None)
    if raw is None or raw == '':
        require(operator != 'between', 'Enter both range values.')
        return lambda value: True
    import math
    try:
        target = float(raw)
        high = float(upper) if operator == 'between' else None
    except (ValueError, TypeError):
        require(False, 'Invalid numeric filter.')
    require(math.isfinite(target) and (high is None or math.isfinite(high) and high >= target), 'Invalid numeric filter range.')
    def matches(value):
        if value is None:
            return False
        return {'gte': lambda: value >= target, 'gt': lambda: value > target,
                'lte': lambda: value <= target, 'lt': lambda: value < target,
                'eq': lambda: value == target, 'between': lambda: target <= value <= high}[operator]()
    return matches


def response_column_filters(raw, test):
    import json
    if raw is None or raw == '':
        return {}
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except (ValueError, TypeError):
            require(False, 'Invalid column filters.')
    allowed = {'driveName', 'name', 'studentId', 'marks', 'minutes', 'shortlisted'}
    allowed |= {'remarks'} if test['type'] == 'aptitude' else {'languageScore', 'reasoningScore'}
    require(isinstance(raw, dict) and set(raw) <= allowed, 'Invalid column filters.')
    output = {}
    for key, value in raw.items():
        require(isinstance(value, dict), 'Invalid column filter.')
        if key in {'marks', 'minutes', 'languageScore', 'reasoningScore'}:
            output[key] = number_filter(value.get('value'), value.get('op', 'gte'), value.get('upper'))
        elif key == 'shortlisted':
            require(isinstance(value.get('values'), list) and all(v in {'Yes', 'No'} for v in value['values']), 'Invalid shortlist filter.')
            output[key] = lambda item, selected=value['values']: item in selected
        else:
            text = value.get('text', '')
            require(isinstance(text, str) and len(text) <= 2000, 'Invalid text filter.')
            output[key] = lambda item, text=text.casefold(): text in str(item or '').casefold()
    return output


def results_rows(test, attempts, minimum_marks=None, minimum_time=None, filters=None, drive_name=None):
    filters = filters or {}
    columns = response_column_filters(filters.get('columnFilters'), test)
    marks_match = number_filter(minimum_marks, filters.get('marksOp', 'gte'), filters.get('maxMarks'))
    time_match = number_filter(minimum_time, filters.get('timeOp', 'gte'), filters.get('maxTime'))
    output = []
    for a in attempts:
        g, elapsed = grade(test['questions'], a), minutes(a)
        if not marks_match(g['total'] if a.get('submittedAt') else None):
            continue
        if not time_match(elapsed):
            continue
        parts = language_parts(test, g) if test['type'] == 'language' and a.get('submittedAt') else [None, None]
        values = {'driveName': drive_name or test.get('name', ''), 'name': a['name'], 'studentId': a['studentId'],
                  'marks': g['total'] if a.get('submittedAt') else None, 'minutes': elapsed,
                  'shortlisted': 'Yes' if a.get('shortlisted') else 'No', 'remarks': a.get('shortlistRemarks', ''),
                  'languageScore': parts[0], 'reasoningScore': parts[1]}
        if not all(matches(values[key]) for key, matches in columns.items()):
            continue
        output.append((a, g, elapsed))
    return sorted(output, key=lambda row: (not bool(row[0].get('submittedAt')), -row[1]['total'], row[2] or 0))


def language_parts(test, grading):
    scores = {result['id']: result['score'] for result in grading['results']}
    values = []
    for group in (test['questions'][:20], test['questions'][20:25]):
        marks = [scores.get(q['id']) for q in group]
        values.append(None if any(mark is None for mark in marks) else round(sum(marks), 6))
    return values


def style_sheet(sheet, widths=None):
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='038681')
    for index in range(1, sheet.max_column + 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = (widths[index - 1] if widths and index <= len(widths) else 24)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions


def workbook_bytes(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def excel(drive, test, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Candidate results'
    language = test['type'] == 'language'
    sheet.append(['Drive name', 'Student name', 'Student ID', 'Total marks' if language else 'Marks obtained'] +
                 (['Language score (20)', 'Reasoning score (5)'] if language else []) +
                 ['Time taken (minutes)', 'Pending questions', 'Shortlisted', 'Assessment', 'Remarks'])
    for a, g, elapsed in rows:
        sheet.append([drive['name'], a['name'], a['studentId'], g['total'] if a.get('submittedAt') else None] +
                     (language_parts(test, g) if language and a.get('submittedAt') else [None, None] if language else []) +
                     [elapsed, g['pending'] if a.get('submittedAt') else None,
                      'Yes' if a.get('shortlisted') else 'No', test['type'], a.get('shortlistRemarks', '')])
    style_sheet(sheet, [34, 28, 20])
    return workbook_bytes(book)


def gd_excel(drive, mode='roster'):
    book = Workbook()
    gd = drive['gd']
    if mode == 'roster':
        sheet = book.active
        sheet.title = 'GD pools'
        sheet.append(['Drive name', 'Pool', '#', 'Candidate name', 'Student ID'])
        for pool in gd['pools']:
            for index, key in enumerate(pool['members'], start=1):
                member = gd['candidates'][key]
                sheet.append([drive['name'], pool['name'], index, member['name'], member['studentId']])
        style_sheet(sheet, [34, 18, 8, 32, 24])
    else:
        book.remove(book.active)
        fields = ('initiative', 'content', 'communication', 'bodyLanguage')
        for index, pool in enumerate(gd['pools'], start=1):
            sheet = book.create_sheet(pool['name'][:31])
            sheet.append(['Drive name', 'Pool', 'Evaluator', '#', 'Candidate name', 'Student ID',
                          'Initiative (5)', 'Content (5)', 'Communication (5)', 'Body Language (5)', 'Total (20)', 'Recommend'])
            for number, key in enumerate(pool['members'], start=1):
                member, score = gd['candidates'][key], gd.get('scores', {}).get(key, {})
                sheet.append([drive['name'], pool['name'], pool.get('evaluatorName', ''), number, member['name'], member['studentId']] +
                             [score.get(field) for field in fields] + [score.get('total'), 'Yes' if score.get('shortlisted') else 'No'])
            style_sheet(sheet, [34, 18, 26, 8, 32, 24])
    return workbook_bytes(book)


def words(drive, test, attempts):
    stream, names = io.BytesIO(), set()
    with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
        for a in attempts:
            doc = Document()
            section = doc.sections[0]
            section.top_margin = section.bottom_margin = Inches(.75)
            doc.styles['Normal'].font.name = 'Calibri'
            doc.styles['Normal'].font.size = Pt(11)
            title = doc.add_paragraph('Prompt Engineering response')
            title.runs[0].bold = True
            doc.add_paragraph(drive['name'])
            doc.add_paragraph(f"Student name: {a['name']}\nStudent ID: {a['studentId']}\nTime taken: {minutes(a):.2f} minutes\nSubmission: {a.get('reason', 'submitted')}")
            for i, q in enumerate(test['questions']):
                p = doc.add_paragraph(f'Question {i + 1}')
                p.runs[0].bold = True
                doc.add_paragraph(q['text'])
                p = doc.add_paragraph('Response')
                p.runs[0].bold = True
                doc.add_paragraph(a.get('responses', {}).get(q['id']) or '(No answer entered)')
            data = io.BytesIO()
            doc.save(data)
            base, suffix = filename(a['name'] + '_' + a['studentId']), 1
            name = base + '.docx'
            while name.casefold() in names:
                suffix += 1
                name = f'{base}_{suffix}.docx'
            names.add(name.casefold())
            archive.writestr(name, data.getvalue())
    return stream.getvalue()



def final_excel(drive, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Final scores'
    sheet.append(['Drive name', 'Candidate name', 'Student ID', 'Aptitude score', 'Shortlisted in Aptitude',
                  'GD score (20)', 'GD recommend', 'GD evaluator', 'Language score (20)', 'Reasoning score (5)', 'Interview shortlisted',
                  'Interview remarks', 'Prompt Engineering score (10)', 'Interviewer', 'Interview comments', 'Final status'])
    for row in rows:
        sheet.append([drive['name'], row['name'], row['studentId'], row.get('aptitudeScore'),
                      'Yes' if row.get('aptitudeShortlisted') else 'No', row.get('gdScore'),
                      'Yes' if row.get('gdRecommended') else 'No', row.get('gdEvaluator', ''), row.get('languageScore'), row.get('reasoningScore'),
                      'Yes' if row.get('interviewShortlisted') else 'No', row.get('interviewRemarks', ''), row.get('promptScore'),
                      row.get('interviewer', ''), row.get('interviewComments', ''),
                      {'none': 'None', 'selected': 'Selected', 'on-hold': 'On hold', 'rejected': 'Rejected'}.get(row.get('outcome'), 'None')])
    style_sheet(sheet, [34, 28, 20, 18, 24, 16, 18, 26, 22, 22, 24, 45, 30, 26, 50, 18])
    return workbook_bytes(book)


def prompt_workbook(drive, test, candidates):
    book = Workbook()
    book.remove(book.active)
    names = set()
    for candidate in candidates:
        base = re.sub(r'[\\/*?:\[\]:]', '_', candidate.get('name', '')).strip().strip("'") or 'Candidate'
        if base.casefold() == 'history':
            base = 'Candidate History'
        title, suffix = base[:31], 1
        while title.casefold() in names:
            suffix += 1
            tail = f' ({suffix})'
            title = base[:31-len(tail)] + tail
        names.add(title.casefold())
        sheet = book.create_sheet(title)
        sheet.append(['Field', 'Details'])
        for label, value in [('Drive', drive['name']), ('Candidate name', candidate['name']),
                             ('Student ID', candidate['studentId']), ('Score (out of 10)', candidate.get('promptScore')),
                             ('Time taken (minutes)', minutes(candidate))]:
            sheet.append([label, value])
        for index, question in enumerate(test['questions'], 1):
            text = question['text']
            for offset in range(0, len(text), 30000):
                sheet.append([f'Question {index}' if offset == 0 else 'Question continued', text[offset:offset + 30000]])
            response = candidate.get('responses', {}).get(question['id'])
            if not response and len(test['questions']) == 1:
                response = candidate.get('answer')
            response = str(response or '(No answer entered)')
            for offset in range(0, len(response), 30000):
                sheet.append(['Response' if offset == 0 else 'Response continued', response[offset:offset + 30000]])
        style_sheet(sheet, [25, 110])
        for row in sheet.iter_rows(min_row=2):
            row[0].font = Font(bold=True, color='28544E')
            row[1].alignment = Alignment(vertical='top', wrap_text=True)
            sheet.row_dimensions[row[0].row].height = min(300, max(28, 15 * (str(row[1].value or '').count('\n') + len(str(row[1].value or '')) // 110 + 1)))
    return workbook_bytes(book)
