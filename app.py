"""TurnB Recruitment Platform. Admin credentials and all assessment data stay server-side."""
import copy
import hashlib
import hmac
import io
import logging
import math
import os
import secrets
import threading
import time
import uuid
from datetime import timedelta
from functools import wraps
from openpyxl import load_workbook

from flask import Flask, jsonify, request, send_file, session
from sqlalchemy import delete, insert, select, update
from sqlalchemy.exc import SQLAlchemyError
from werkzeug.exceptions import HTTPException
from werkzeug.security import check_password_hash

from configuration import update_configuration
from domain import (DURATIONS, RuleError, clean_text, grade, require,
                    public_questions, responses_validated, student_key)
from exports import excel, final_excel, filename, results_rows, words, gd_excel, language_parts, prompt_workbook
from store import Store, attempts, login_limits, now_ms


def create_app(settings=None):
    app = Flask(__name__, static_folder='static')
    app.config.update(SECRET_KEY=os.environ.get('SESSION_SECRET') or secrets.token_urlsafe(48),
                      ADMIN_EMAIL=os.environ.get('ADMIN_EMAIL', 'turnbrecruitment@gmail.com'),
                      ADMIN_PASSWORD_HASH=os.environ.get('ADMIN_PASSWORD_HASH', ''),
                      DATABASE_URL=os.environ.get('DATABASE_URL', ''),
                      PUBLIC_ORIGIN=os.environ.get('PUBLIC_ORIGIN', ''),
                      APP_ENV=os.environ.get('APP_ENV', 'production'),
                      CONFIGURED_SECRET=bool(os.environ.get('SESSION_SECRET')),
                      SESSION_COOKIE_NAME='turnb_admin', SESSION_COOKIE_HTTPONLY=True,
                      SESSION_COOKIE_SECURE=True, SESSION_COOKIE_SAMESITE='Strict',
                      PERMANENT_SESSION_LIFETIME=timedelta(hours=8), MAX_CONTENT_LENGTH=8 * 1024 * 1024)
    if settings:
        app.config.update(settings)
    local = app.config['APP_ENV'] in {'test', 'development'}
    if local:
        app.config['SESSION_COOKIE_SECURE'] = False
    store = None
    if app.config['DATABASE_URL']:
        store = Store(app.config['DATABASE_URL'], testing=local)
        if local:
            store.initialize()
    app.extensions['store'] = store

    def ready():
        return store is not None and (local or app.config['CONFIGURED_SECRET'] and len(app.secret_key) >= 32) and bool(app.config['ADMIN_PASSWORD_HASH'])

    def db_required():
        if not ready():
            return jsonify(error='The application requires database and administrator configuration.'), 503
        return None

    def admin(function):
        @wraps(function)
        def wrapped(*args, **kwargs):
            if not session.get('admin'):
                return jsonify(error='Administrator login required.'), 401
            return function(*args, **kwargs)
        return wrapped

    @app.before_request
    def security():
        if request.path.startswith('/api/'):
            problem = db_required()
            if problem:
                return problem
            if request.method not in {'GET', 'HEAD', 'OPTIONS'}:
                is_aptitude_import = (request.path.startswith('/api/admin/drives/') and
                                      request.path.endswith(('/aptitude/import', '/aptitude/responses-import')))
                if request.mimetype != 'application/json' and not (is_aptitude_import and request.mimetype == 'multipart/form-data'):
                    return jsonify(error='JSON request required.'), 415
                expected = session.get('csrf')
                if not expected or not hmac.compare_digest(request.headers.get('X-CSRF-Token', ''), expected):
                    return jsonify(error='Refresh the page and try again.'), 403
                origin = request.headers.get('Origin')
                allowed = app.config['PUBLIC_ORIGIN'] or request.host_url.rstrip('/')
                if origin and origin != allowed:
                    return jsonify(error='Origin not allowed.'), 403

    @app.after_request
    def response_headers(response):
        response.headers.update({'Cache-Control': 'no-store', 'X-Content-Type-Options': 'nosniff',
                                 'X-Frame-Options': 'DENY', 'Referrer-Policy': 'no-referrer',
                                 'Permissions-Policy': 'camera=(), microphone=(), geolocation=()',
                                 'Content-Security-Policy': "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'; object-src 'none'"})
        if not local:
            response.headers['Strict-Transport-Security'] = 'max-age=31536000'
        return response

    @app.errorhandler(RuleError)
    def invalid(error):
        return jsonify(error=str(error)), 400

    @app.errorhandler(SQLAlchemyError)
    def database_error(error):
        app.logger.error('Database operation failed: %s', type(error).__name__)
        return jsonify(error='Storage is temporarily unavailable. Please retry; do not close your test.'), 503

    @app.errorhandler(Exception)
    def unexpected(error):
        if isinstance(error, HTTPException):
            return jsonify(error=error.description), error.code
        app.logger.error('Request failed: %s', type(error).__name__)
        return jsonify(error='The request could not be completed. Please retry.'), 500

    @app.get('/')
    @app.get('/test/<test_id>')
    def index(test_id=None):
        return app.send_static_file('index.html')

    @app.get('/healthz')
    def health():
        return jsonify(status='ok', service='turnb-recruitment')

    @app.get('/readyz')
    def readiness():
        if not ready():
            return jsonify(status='configuration-required'), 503
        with store.transaction():
            pass
        return jsonify(status='ready')

    @app.get('/api/auth/session')
    def auth_session():
        session.setdefault('csrf', secrets.token_urlsafe(32))
        return jsonify(admin=bool(session.get('admin')), csrf=session['csrf'])

    @app.post('/api/auth/login')
    def login():
        data = request.get_json()
        require(isinstance(data, dict), 'Invalid login.')
        email = clean_text(data.get('email', ''), 200, 'email').casefold()
        password = data.get('password', '')
        require(isinstance(password, str) and len(password) <= 1024, 'Invalid password.')
        # Azure terminates ingress; use the last forwarded address, not a user-supplied first one.
        address = request.headers.get('X-Forwarded-For', request.remote_addr or '').split(',')[-1].strip()
        bucket = hmac.new(app.secret_key.encode(), address.encode(), hashlib.sha256).hexdigest()
        timestamp = now_ms()
        with store.transaction() as tx:
            tx.connection.execute(delete(login_limits).where(login_limits.c.window < timestamp - 900000))
            row = tx.connection.execute(select(login_limits).where(login_limits.c.id == bucket)).mappings().first()
            if row and row['count'] >= 5:
                return jsonify(error='Too many login attempts. Try again in 15 minutes.'), 429
            if row:
                tx.connection.execute(update(login_limits).where(login_limits.c.id == bucket).values(count=row['count'] + 1))
            else:
                tx.connection.execute(insert(login_limits).values(id=bucket, window=timestamp, count=1))
        valid_hash = check_password_hash(app.config['ADMIN_PASSWORD_HASH'], password)
        if not (valid_hash and hmac.compare_digest(email, app.config['ADMIN_EMAIL'].casefold())):
            return jsonify(error='Email or password is incorrect.'), 401
        with store.transaction() as tx:
            tx.connection.execute(delete(login_limits).where(login_limits.c.id == bucket))
        session.clear()
        session.update(admin=True, csrf=secrets.token_urlsafe(32))
        session.permanent = True
        return jsonify(admin=True, csrf=session['csrf'])

    @app.post('/api/auth/logout')
    def logout():
        session.clear()
        return jsonify(ok=True)

    @app.get('/api/admin/state')
    @admin
    def get_state():
        with store.transaction() as tx:
            tx.expire()
            return jsonify(state=tx.export_state(), revision=tx.next_version(), serverNow=now_ms())

    @app.put('/api/admin/state')
    @admin
    def put_state():
        payload = request.get_json()
        require(isinstance(payload, dict), 'Invalid request.')
        with store.transaction() as tx:
            tx.expire()
            if payload.get('revision') != tx.next_version():
                return jsonify(error='Another change was saved. Reload before continuing.'), 409
            update_configuration(tx, payload.get('state'))
            return jsonify(revision=tx.next_version(), state=tx.export_state(), serverNow=now_ms())

    @app.post('/api/admin/drives/<drive_id>/members')
    @admin
    def add_drive_members(drive_id):
        data = request.get_json()
        require(isinstance(data, dict) and isinstance(data.get('members'), list), 'Provide member names.')
        require(len(data['members']) <= 100, 'Add no more than 100 drive members.')
        additions = [clean_text(value, 100, 'drive member') for value in data['members']]
        require(additions and all(additions), 'Enter at least one member name.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None, 'Drive not found.')
            members = list(drive.get('members', []))
            seen = {student_key(name) for name in members}
            for name in additions:
                if student_key(name) not in seen:
                    members.append(name)
                    seen.add(student_key(name))
            require(len(members) <= 100, 'Add no more than 100 drive members.')
            drive['members'] = members
            return jsonify(members=members, serverNow=now_ms())

    def test_required(tx, test_id):
        test = tx.test(test_id)
        require(test is not None, 'This test link is unavailable.')
        return test

    @app.post('/api/admin/drives/<drive_id>/aptitude/import')
    @admin
    def import_aptitude_results(drive_id):
        uploaded = request.files.get('file')
        require(uploaded is not None and uploaded.filename.lower().endswith('.xlsx'),
                'Choose the Aptitude score tracker as an .xlsx file.')
        try:
            workbook = load_workbook(io.BytesIO(uploaded.read()), read_only=True, data_only=True)
            sheet = next((ws for ws in workbook.worksheets if ws.title.strip().casefold() == 'shortlist'), None)
            require(sheet is not None, 'The workbook must contain a sheet named Shortlist.')
            headers = [str(value or '').strip().casefold() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
            expected = ['full name', 'student id', 'marks obtained', 'time taken (minutes)']
            require(headers[:4] == expected, 'Shortlist columns must be Full Name, Student ID, Marks Obtained, and Time Taken (Minutes), in that order.')
            records, seen = [], set()
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                values = list(row[:4])
                if not any(value not in (None, '') for value in values):
                    continue
                require(len(values) == 4 and all(value not in (None, '') for value in values),
                        f'Complete all four Shortlist fields on row {row_number}.')
                name = clean_text(str(values[0]), 100, 'candidate name')
                student_id = clean_text(str(values[1]), 60, 'student ID')
                require(name and student_id, f'Name and student ID are required on row {row_number}.')
                try:
                    score, elapsed = float(values[2]), float(values[3])
                except (TypeError, ValueError):
                    raise RuleError(f'Marks and time must be numeric on row {row_number}.')
                require(math.isfinite(score) and 0 <= score <= 30, f'Marks must be between 0 and 30 on row {row_number}.')
                require(math.isfinite(elapsed) and 0 <= elapsed <= 1440, f'Time must be between 0 and 1440 minutes on row {row_number}.')
                key = student_key(student_id)
                require(key not in seen, f'Duplicate student ID on row {row_number}.')
                seen.add(key)
                records.append({'name': name, 'studentId': student_id,
                                'importedScore': round(score, 4),
                                'importedElapsedMinutes': round(elapsed, 4)})
                require(len(records) <= 2000, 'The Shortlist sheet exceeds the 2,000 candidate limit.')
            workbook.close()
        except RuleError:
            raise
        except Exception:
            raise RuleError('The workbook could not be read. Save it as a valid .xlsx file and try again.')
        require(records, 'The Shortlist sheet contains no candidate rows.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None, 'Drive not found.')
            require(not drive.get('finishedAt'), 'Cannot import results into a finished drive.')
            require(not any(s['driveId'] == drive_id and s['type'] == 'aptitude' for s in tx.state['sessions']),
                    'This drive already has an Aptitude session. Import is allowed only once.')
            now = now_ms()
            test_id = str(uuid.uuid4())
            test = {'id': test_id, 'driveId': drive_id, 'type': 'aptitude',
                    'name': drive['name'] + ' · Aptitude', 'start': now, 'closedAt': now,
                    'endReason': 'imported', 'historicalImport': True, 'questions': []}
            tx.state['sessions'].insert(0, test)
            for record in records:
                elapsed_ms = int(record['importedElapsedMinutes'] * 60000)
                started = now - elapsed_ms
                attempt = {'id': str(uuid.uuid4()), **record, 'startedAt': started,
                           'deadline': started + 15 * 60000, 'submittedAt': now,
                           'reason': 'imported', 'responses': {}, 'sequence': 0,
                           'manualMarks': {}, 'shortlisted': True,
                           'importedSource': 'Shortlist'}
                key_hash = hashlib.sha256(student_key(record['studentId']).encode()).hexdigest()
                token_hash = hashlib.sha256(secrets.token_bytes(32)).hexdigest()
                tx.add_attempt(test_id, attempt, key_hash, token_hash)
            return jsonify(imported=len(records), testId=test_id, serverNow=now)

    @app.post('/api/admin/drives/<drive_id>/aptitude/responses-import')
    @admin
    def import_aptitude_responses(drive_id):
        from response_import import read_responses, preview_fingerprint
        uploaded = request.files.get('file')
        require(uploaded is not None and uploaded.filename.lower().endswith('.xlsx'), 'Choose an .xlsx response workbook.')
        raw = uploaded.read()
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None, 'Drive not found.')
            require(not drive.get('finishedAt'), 'Cannot import into a finished drive.')
            test = next((s for s in tx.state['sessions'] if s['driveId'] == drive_id and s['type'] == 'aptitude'), None)
            existing = tx.all_attempts(test['id']) if test else []
            require(not test or test.get('historicalImport') or not existing, 'This drive has online Aptitude attempts; importing would overwrite live test data.')
            questions = copy.deepcopy(test.get('questions') if test and test.get('questions') else tx.state['questionSets']['aptitude'])
            records, errors, mapping = read_responses(raw, questions)
            fingerprint = preview_fingerprint(raw, questions, (tx.version, existing))
            by_key = {student_key(a['studentId']): a for a in existing}
            preview = [{'name': r['name'], 'studentId': r['studentId'], 'score': grade(questions, r)['total'],
                        'minutes': r['importedElapsedMinutes'], 'action': 'Update' if student_key(r['studentId']) in by_key else 'Add'} for r in records]
            if request.form.get('confirm') != 'true':
                return jsonify(records=preview, errors=errors, mapping=mapping, fingerprint=fingerprint,
                               maxScore=sum(q['correctScore'] for q in questions), serverNow=now_ms())
            require(not errors, 'Resolve all worksheet errors before importing.')
            require(hmac.compare_digest(request.form.get('fingerprint', ''), fingerprint), 'Drive data or answer keys changed. Preview the file again.')
            now = now_ms()
            if not test:
                test = {'id': str(uuid.uuid4()), 'driveId': drive_id, 'type': 'aptitude',
                        'name': drive['name'] + ' · Aptitude', 'start': now, 'closedAt': now,
                        'durationMinutes': tx.state.get('durations', {}).get('aptitude', 15), 'questions': questions}
                tx.state['sessions'].insert(0, test)
            test.update(questions=questions, historicalImport=True, responseImport=True, closedAt=now, endReason='imported')
            for record in records:
                old = by_key.get(student_key(record['studentId']))
                started = now - int(record['importedElapsedMinutes'] * 60000)
                if old is not None:
                    # Preserve identity, shortlist decisions and all later-round data.
                    old.update(responses=record['responses'], importedElapsedMinutes=record['importedElapsedMinutes'],
                               manualMarks={}, questionOrder=[q['id'] for q in questions], optionOrder={},
                               startedAt=started, submittedAt=now, deadline=now, reason='imported', importedSource='Input Data')
                    old.pop('importedScore', None)
                else:
                    attempt = {'id': str(uuid.uuid4()), **record, 'startedAt': started, 'submittedAt': now,
                               'deadline': now, 'reason': 'imported', 'sequence': 0, 'manualMarks': {},
                               'shortlisted': False, 'importedSource': 'Input Data',
                               'questionOrder': [q['id'] for q in questions], 'optionOrder': {}}
                    tx.add_attempt(test['id'], attempt, hashlib.sha256(student_key(record['studentId']).encode()).hexdigest(),
                                   hashlib.sha256(secrets.token_bytes(32)).hexdigest())
            return jsonify(imported=len(records), testId=test['id'], serverNow=now)

    def candidate_attempt(tx, test):
        token = request.cookies.get('attempt_' + test['id'], '')
        require(token and len(token) <= 200, 'Start the test in this browser first.')
        digest = hashlib.sha256(token.encode()).hexdigest()
        for a in tx.all_attempts(test['id']):
            if hmac.compare_digest(tx.loaded[a['id']]['row']['token_hash'], digest):
                return a
        raise RuleError('Attempt credentials are invalid.')

    def matching_token(test_id, attempt_id, qid, index):
        value = f'{test_id}:{attempt_id}:{qid}:{index}'
        return hmac.new(app.secret_key.encode(), value.encode(), hashlib.sha256).hexdigest()[:32]

    def candidate_payload(test, attempt=None):
        public = {k: test[k] for k in ('id', 'type', 'name', 'driveId', 'closedAt')}
        public['availableAt'] = test.get('availableAt')
        public['durationMinutes'] = test.get('durationMinutes', DURATIONS[test['type']])
        public['questionCount'] = len(test['questions'])
        public['scoringRules'] = {'correct': test['questions'][0].get('correctScore', 1) if test['type'] != 'prompt' and test.get('questions') else 1,
                                  'wrong': test['questions'][0].get('wrongScore', 0) if test['type'] != 'prompt' and test.get('questions') else 0}
        if attempt is not None:
            signer = lambda qid, i: matching_token(test['id'], attempt['id'], qid, i)
            by_id = {q['id']: q for q in test['questions']}
            order = attempt.get('questionOrder') or [q['id'] for q in test['questions']]
            ordered = [by_id[qid] for qid in order if qid in by_id]
            ordered.extend(q for q in test['questions'] if q['id'] not in set(order))
            public['questions'] = public_questions(ordered, signer)
            option_order = attempt.get('optionOrder', {})
            for q in public['questions']:
                permutation = option_order.get(q['id'])
                if permutation and q.get('options') and len(permutation) == len(q['options']):
                    q['options'] = [q['options'][i] for i in permutation]
            candidate = {k: copy.deepcopy(v) for k, v in attempt.items()
                         if k not in {'manualMarks', 'shortlisted', 'manualShortlisted', 'shortlistRemarks', 'promptScore', 'questionOrder', 'optionOrder'}}
            for q in test['questions']:
                if q['type'] == 'matching' and q['id'] in candidate['responses']:
                    candidate['responses'][q['id']] = [None if i is None else signer(q['id'], i) for i in candidate['responses'][q['id']]]
                elif q.get('options') and q['id'] in candidate['responses']:
                    permutation = option_order.get(q['id']) or list(range(len(q['options'])))
                    display_index = {original: display for display, original in enumerate(permutation)}
                    response = candidate['responses'][q['id']]
                    if isinstance(response, list):
                        candidate['responses'][q['id']] = [display_index[i] for i in response if type(i) is int and i in display_index]
            public['attempts'] = [candidate]
        else:
            public['questions'], public['attempts'] = [], []
        return {'test': public, 'serverNow': now_ms()}

    def aptitude_identity(tx, drive_id, student_id):
        key = student_key(student_id)
        aptitude_tests = [s for s in tx.state['sessions'] if s['driveId'] == drive_id and s['type'] == 'aptitude']
        for aptitude in aptitude_tests:
            for candidate in tx.all_attempts(aptitude['id']):
                if student_key(candidate['studentId']) == key:
                    return candidate
        return None

    @app.post('/api/candidate/<test_id>/identity')
    def candidate_identity(test_id):
        data = request.get_json()
        require(isinstance(data, dict), 'Invalid candidate details.')
        student = clean_text(data.get('studentId'), 60, 'student ID')
        require(bool(student), 'Student ID is required.')
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            require(not test.get('availableAt') or now_ms() >= test['availableAt'], 'This test has not started yet. Please return at the scheduled time.')
            require(test['type'] in {'language', 'prompt'}, 'Enter your name and student ID for Aptitude.')
            require(not test.get('closedAt'), 'This test has ended.')
            candidate = aptitude_identity(tx, test['driveId'], student)
            require(candidate is not None, 'Student ID not found. Please enter the correct student ID used for Aptitude.')
            return jsonify(name=candidate['name'], studentId=candidate['studentId'])

    @app.get('/api/candidate/<test_id>')
    def get_candidate(test_id):
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            tx.expire(test_id)
            a = candidate_attempt(tx, test) if request.cookies.get('attempt_' + test_id) else None
            return jsonify(candidate_payload(test, a))

    @app.post('/api/candidate/<test_id>/start')
    def start_candidate(test_id):
        data = request.get_json()
        require(isinstance(data, dict), 'Invalid candidate details.')
        student = clean_text(data.get('studentId'), 60, 'student ID')
        require(bool(student), 'Student ID is required.')
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            require(not test.get('availableAt') or now_ms() >= test['availableAt'], 'This test has not started yet. Please return at the scheduled time.')
            if test['type'] == 'aptitude':
                name = clean_text(data.get('name'), 100, 'name')
                require(bool(name), 'Name is required.')
            else:
                identity = aptitude_identity(tx, test['driveId'], student)
                require(identity is not None, 'Student ID not found. Please enter the correct student ID used for Aptitude.')
                name, student = identity['name'], identity['studentId']
            tx.expire(test_id)
            require(not test.get('closedAt'), 'This test has ended.')
            if request.cookies.get('attempt_' + test_id):
                return jsonify(candidate_payload(test, candidate_attempt(tx, test)))
            require(test['questions'], 'Questions are not configured yet.')
            key = student_key(student)
            require(not any(student_key(a['studentId']) == key for a in tx.all_attempts(test_id)), 'This student ID already has an attempt. Use the original browser to continue.')
            token, now = secrets.token_urlsafe(48), now_ms()
            duration = test.get('durationMinutes', DURATIONS[test['type']])
            question_order = [q['id'] for q in test['questions']]
            option_order = {}
            rng = secrets.SystemRandom()
            if test['type'] == 'aptitude':
                rng.shuffle(question_order)
                shuffle_options = set(question_order)
            elif test['type'] == 'language':
                language_count = min(20, len(question_order))
                first_questions = question_order[:language_count]
                rng.shuffle(first_questions)
                question_order = first_questions + question_order[language_count:]
                shuffle_options = set(first_questions)
            else:
                shuffle_options = set()
            for q in test['questions']:
                if q['id'] in shuffle_options and q.get('options') and q['type'] in {'single', 'multiple', 'boolean', 'ranking'}:
                    permutation = list(range(len(q['options'])))
                    rng.shuffle(permutation)
                    option_order[q['id']] = permutation
            a = {'id': str(uuid.uuid4()), 'name': name, 'studentId': student, 'startedAt': now,
                 'deadline': now + duration * 60000, 'responses': {}, 'sequence': 0,
                 'submittedAt': None, 'reason': None, 'manualMarks': {}, 'shortlisted': False,
                 'questionOrder': question_order, 'optionOrder': option_order}
            tx.add_attempt(test_id, a, hashlib.sha256(key.encode()).hexdigest(), hashlib.sha256(token.encode()).hexdigest())
            response = jsonify(candidate_payload(test, a))
            response.set_cookie('attempt_' + test_id, token, httponly=True, secure=not local,
                                samesite='Strict', path='/api/candidate/' + test_id, max_age=86400)
            return response

    @app.post('/api/candidate/<test_id>/answer')
    def candidate_answer(test_id):
        data = request.get_json()
        require(isinstance(data, dict), 'Invalid answer.')
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            a = candidate_attempt(tx, test)
            tx.expire(test_id)
            if not a.get('submittedAt') and not test.get('closedAt'):
                sequence = data.get('sequence')
                require(type(sequence) is int and 0 < sequence < 2**31, 'Invalid save sequence.')
                if sequence > a['sequence']:
                    responses = copy.deepcopy(data.get('responses', {}))
                    require(isinstance(responses, dict), 'Invalid response object.')
                    option_order = a.get('optionOrder', {})
                    for q in test['questions']:
                        if q['type'] == 'matching' and q['id'] in responses:
                            tokens = {matching_token(test['id'], a['id'], q['id'], i): i for i in range(len(q['pairs']))}
                            values = responses[q['id']]
                            require(isinstance(values, list) and all(v is None or isinstance(v, str) and v in tokens for v in values), 'Invalid matching choice.')
                            responses[q['id']] = [None if v is None else tokens[v] for v in values]
                        elif q.get('options') and q['id'] in responses:
                            permutation = option_order.get(q['id']) or list(range(len(q['options'])))
                            values = responses[q['id']]
                            require(isinstance(values, list) and all(type(v) is int and 0 <= v < len(permutation) for v in values), 'Invalid option choice.')
                            responses[q['id']] = [permutation[v] for v in values]
                    a['responses'] = responses_validated(test['questions'], responses)
                    a['sequence'] = sequence
                if data.get('submit') is True:
                    a.update(submittedAt=now_ms(), reason='manual')
                tx.expire(test_id)
            return jsonify(candidate_payload(test, a))

    def filter_numbers(data):
        values = []
        for field in ('minMarks', 'minTime'):
            value = data.get(field)
            if value is None or value == '':
                values.append(None)
                continue
            try:
                number = float(value)
            except (ValueError, TypeError):
                raise RuleError('Invalid numeric filter.')
            require(math.isfinite(number) and (field != 'minTime' or number >= 0), 'Invalid filter range.')
            values.append(number)
        return values

    @app.post('/api/admin/tests/<test_id>/shortlist')
    @admin
    def shortlist(test_id):
        data = request.get_json()
        require(isinstance(data, dict), 'Invalid request.')
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            require(test['type'] != 'prompt', 'Prompt Engineering is evaluated separately.')
            tx.expire(test_id)
            rows = results_rows(test, tx.all_attempts(test_id), *filter_numbers(data), filters=data, drive_name=next(d['name'] for d in tx.state['drives'] if d['id'] == test['driveId']))
            changed, skipped = 0, 0
            for a, g, _ in rows:
                if data.get('attemptId') and a['id'] != data['attemptId']:
                    continue
                if not a.get('submittedAt') or g['pending']:
                    skipped += 1
                    continue
                a['shortlisted'] = data.get('value', True) is True
                a['manualShortlisted'] = False
                a['shortlistRemarks'] = ''
                changed += 1
            if test['type'] == 'language' and data.get('value', True) is True and changed:
                test['shortlistAppliedAt'] = now_ms()
            return jsonify(changed=changed, skipped=skipped)

    @app.post('/api/admin/tests/<test_id>/manual-shortlist')
    @admin
    def manual_aptitude_shortlist(test_id):
        data = request.get_json()
        require(isinstance(data, dict) and isinstance(data.get('changes'), list) and 0 < len(data['changes']) <= 2000,
                'Provide the candidate updates.')
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            require(test['type'] == 'aptitude', 'Manual aptitude shortlisting is available only for Aptitude.')
            tx.expire(test_id)
            candidates = {a['id']: a for a in tx.all_attempts(test_id)}
            validated, seen = [], set()
            for item in data['changes']:
                require(isinstance(item, dict), 'Invalid candidate update.')
                aid = item.get('attemptId')
                require(isinstance(aid, str) and aid in candidates and aid not in seen, 'Candidate not found or repeated.')
                seen.add(aid)
                a = candidates[aid]
                value = item.get('shortlisted')
                require(type(value) is bool, 'Choose a valid shortlist status.')
                remarks = clean_text(item.get('remarks', ''), 2000, 'shortlist remarks')
                require(not value or bool(remarks), 'Remarks are required for every manually shortlisted candidate.')
                require(a.get('submittedAt') and not grade(test['questions'], a)['pending'],
                        'Only completed candidates with final marks can be manually shortlisted.')
                validated.append((a, value, remarks))
            for a, value, remarks in validated:
                if value:
                    a.update(shortlisted=True, manualShortlisted=True, shortlistRemarks=remarks)
                elif a.get('manualShortlisted'):
                    a.update(shortlisted=False, manualShortlisted=False, shortlistRemarks='')
            return jsonify(changed=len(validated), serverNow=now_ms())

    @app.get('/api/admin/drives/<drive_id>/gd/export')
    @admin
    def export_gd_pools(drive_id):
        mode = request.args.get('mode', 'roster')
        require(mode in {'roster', 'scores'}, 'Choose a valid GD export.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None and drive.get('gd', {}).get('pools'), 'Generate GD pools before exporting.')
            payload = gd_excel(drive, mode)
            name = filename(drive['name'] + '_gd_' + mode) + '.xlsx'
        return send_file(io.BytesIO(payload), as_attachment=True, download_name=name,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.post('/api/admin/drives/<drive_id>/delete')
    @admin
    def delete_drive(drive_id):
        data = request.get_json()
        require(isinstance(data, dict) and data.get('confirm') is True, 'Confirm drive deletion.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None, 'Drive not found.')
            test_ids = [s['id'] for s in tx.state['sessions'] if s['driveId'] == drive_id]
            if test_ids:
                tx.connection.execute(delete(attempts).where(attempts.c.test_id.in_(test_ids)))
            tx.state['sessions'] = [s for s in tx.state['sessions'] if s['driveId'] != drive_id]
            tx.state['drives'] = [d for d in tx.state['drives'] if d['id'] != drive_id]
            return jsonify(deleted=True, tests=len(test_ids))

    @app.post('/api/admin/drives/<drive_id>/gd/pools')
    @admin
    def generate_gd_pools(drive_id):
        data = request.get_json()
        require(isinstance(data, dict) and type(data.get('poolSize')) is int and 1 <= data['poolSize'] <= 10, 'Pool size must be between 1 and 10.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None, 'Drive not found.')
            require(not drive.get('finishedAt'), 'Cannot generate GD pools for a finished drive.')
            drive_tests = [s for s in tx.state['sessions'] if s['driveId'] == drive_id]
            language = next((s for s in drive_tests if s['type'] == 'language' and
                             (s.get('shortlistAppliedAt') or any(a.get('shortlisted') for a in tx.all_attempts(s['id'])))), None)
            source = language or next((s for s in drive_tests if s['type'] == 'aptitude'), None)
            require(source is not None, 'Add Aptitude results before creating GD pools.')
            tx.expire(source['id'])
            selected = {}
            for candidate in tx.all_attempts(source['id']):
                if candidate.get('shortlisted'):
                    key = student_key(candidate['studentId'])
                    prior = selected.get(key)
                    if prior is None or candidate.get('startedAt', 0) > prior.get('startedAt', 0):
                        selected[key] = {'studentId': candidate['studentId'], 'name': candidate['name']}
            roster = list(selected.items())
            secrets.SystemRandom().shuffle(roster)
            require(bool(roster), 'No shortlisted candidates are available for GD pools.')
            pool_count = max(1, len(roster) // data['poolSize'])
            while math.ceil(len(roster) / pool_count) > 10:
                pool_count += 1
            base, remainder = divmod(len(roster), pool_count)
            pools, offset = [], 0
            for index in range(pool_count):
                size = base + (1 if index < remainder else 0)
                members = [key for key, _ in roster[offset:offset + size]]
                offset += size
                name = f'Pool {chr(65 + index)}' if index < 26 else f'Pool {index + 1}'
                previous_pools = (drive.get('gd') or {}).get('pools', [])
                previous_pool = next((p for p in previous_pools if p.get('name') == name and p.get('members') == members), {})
                pools.append({'name': name, 'members': members, 'evaluatorName': previous_pool.get('evaluatorName', '')})
            previous = drive.get('gd') or {}
            drive['gd'] = {'poolSize': data['poolSize'], 'source': source['type'], 'pools': pools, 'candidates': {key: person for key, person in roster}, 'scores': {key: score for key, score in previous.get('scores', {}).items() if key in selected}, 'updatedAt': now_ms()}
            return jsonify(gd=drive['gd'], candidateCount=len(roster))

    @app.post('/api/admin/drives/<drive_id>/gd/clear')
    @admin
    def clear_gd_pools(drive_id):
        data = request.get_json()
        require(isinstance(data, dict) and data.get('confirm') is True, 'Confirm clearing GD pools.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None and isinstance(drive.get('gd'), dict), 'GD pools not found.')
            require(not drive.get('finishedAt'), 'Cannot clear GD pools for a finished drive.')
            drive.pop('gd')
            return jsonify(cleared=True)

    @app.post('/api/admin/drives/<drive_id>/gd/evaluator')
    @admin
    def save_gd_evaluator(drive_id):
        data = request.get_json()
        require(isinstance(data, dict) and type(data.get('poolIndex')) is int, 'Invalid evaluator request.')
        name = clean_text(data.get('evaluatorName'), 100, 'evaluator name')
        require(bool(name), "Enter the evaluator name.")
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            gd = drive.get('gd') if drive else None
            require(isinstance(gd, dict) and 0 <= data['poolIndex'] < len(gd.get('pools', [])), 'GD pool not found.')
            pool = gd['pools'][data['poolIndex']]
            members = drive.get('members', [])
            if members:
                canonical = next((member for member in members if student_key(member) == student_key(name)), None)
                require(canonical is not None, 'Choose an evaluator from this drive’s members.')
                name = canonical
            pool['evaluatorName'] = name
            gd['updatedAt'] = now_ms()
            return jsonify(pool=pool, serverNow=now_ms())

    @app.post('/api/admin/drives/<drive_id>/gd/score')
    @admin
    def save_gd_score(drive_id):
        data = request.get_json()
        require(isinstance(data, dict), 'Invalid score request.')
        field = data.get('field')
        if field == 'shortlisted':
            value = data.get('value')
            require(type(value) is bool, 'GD shortlist must be true or false.')
        else:
            require(field in {'initiative', 'content', 'communication', 'bodyLanguage'}, 'Invalid GD scoring area.')
            value = data.get('score')
            require(value is None or (type(value) in (int, float) and math.isfinite(value) and 0 <= value <= 5), 'Enter a score from 0 to 5.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None and isinstance(drive.get('gd'), dict), 'Generate GD pools first.')
            key = data.get('studentKey')
            assigned_pool = next((pool for pool in drive['gd']['pools'] if key in pool['members']), None)
            require(assigned_pool is not None, "Candidate is not in this drive GD pools.")
            require(bool(assigned_pool.get('evaluatorName')), 'Save the pool evaluator name before entering scores or shortlisting candidates.')
            score = drive['gd'].setdefault('scores', {}).setdefault(key, {})
            if field == 'shortlisted':
                score['shortlisted'] = value
            elif value is None:
                score.pop(field, None)
            else:
                score[field] = float(value)
            score['total'] = round(sum(score.get(name, 0) for name in ('initiative', 'content', 'communication', 'bodyLanguage')), 2)
            drive['gd']['updatedAt'] = now_ms()
            return jsonify(score=score, serverNow=now_ms())

    @app.post('/api/admin/tests/<test_id>/prompt-score')
    @admin
    def prompt_score(test_id):
        data = request.get_json()
        require(isinstance(data, dict), 'Invalid score request.')
        value = data.get('score')
        require(value is None or (type(value) in (int, float) and math.isfinite(value) and 0 <= value <= 10),
                'Prompt Engineering scores must be from 0 to 10.')
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            require(test['type'] == 'prompt', 'This is not a Prompt Engineering test.')
            attempt = next((a for a in tx.all_attempts(test_id) if a['id'] == data.get('attemptId')), None)
            require(attempt is not None and attempt.get('submittedAt'), 'Score only submitted Prompt Engineering responses.')
            attempt['promptScore'] = None if value is None else round(float(value), 2)
            return jsonify(promptScore=attempt['promptScore'], serverNow=now_ms())

    def final_score_rows(tx, drive_id):
        drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
        require(drive is not None, 'Drive not found.')
        candidates = {}
        for key, member in (drive.get('gd', {}).get('candidates', {}) or {}).items():
            scores = drive.get('gd', {}).get('scores', {}).get(key, {})
            candidates[key] = {'studentKey': key, 'name': member['name'], 'studentId': member['studentId'],
                               'aptitudeScore': None, 'aptitudeShortlisted': False,
                               'gdScore': scores.get('total'), 'gdRecommended': scores.get('shortlisted') is True,
                               'languageScore': None, 'reasoningScore': None, 'promptScore': None,
                               'interviewShortlisted': False, 'interviewManual': False, 'outcome': 'none',
                               '_latest': {}}
        decisions = drive.get('finalDecisions', {})
        for test in (s for s in tx.state['sessions'] if s['driveId'] == drive_id):
            for attempt in tx.all_attempts(test['id']):
                key = student_key(attempt['studentId'])
                row = candidates.setdefault(key, {'studentKey': key, 'name': attempt['name'], 'studentId': attempt['studentId'],
                            'aptitudeScore': None, 'aptitudeShortlisted': False, 'gdScore': None, 'gdRecommended': False,
                            'languageScore': None, 'reasoningScore': None, 'promptScore': None, 'interviewShortlisted': False, 'interviewManual': False, 'outcome': 'none', '_latest': {}})
                if attempt.get('startedAt', 0) >= row.get('_latest_name', -1):
                    row['name'], row['studentId'], row['_latest_name'] = attempt['name'], attempt['studentId'], attempt.get('startedAt', 0)
                if test['type'] not in {'aptitude', 'language', 'prompt'}:
                    continue
                previous = row['_latest'].get(test['type'])
                stamp = (attempt.get('startedAt', 0), attempt['id'])
                if previous and stamp <= previous[0]:
                    continue
                row['_latest'][test['type']] = (stamp, attempt)
                if test['type'] == 'aptitude':
                    row['aptitudeScore'] = grade(test['questions'], attempt)['total'] if attempt.get('submittedAt') else None
                    row['aptitudeShortlisted'] = attempt.get('shortlisted') is True
                elif test['type'] == 'language':
                    row['languageScore'], row['reasoningScore'] = language_parts(test, grade(test['questions'], attempt)) if attempt.get('submittedAt') else (None, None)
                else:
                    row['promptScore'] = attempt.get('promptScore')
        for key, row in candidates.items():
            decision = decisions.get(key, {})
            row['interviewShortlisted'] = decision.get('interviewShortlisted') is True
            row['interviewManual'] = decision.get('interviewManual') is True
            row['outcome'] = decision.get('outcome') or 'none'
            row['interviewRemarks'] = decision.get('interviewRemarks', '')
            row['interviewer'] = decision.get('interviewer', '')
            row['gdEvaluator'] = next((p.get('evaluatorName', '') for p in drive.get('gd', {}).get('pools', []) if key in p.get('members', [])), '')
            row.pop('_latest', None)
            row.pop('_latest_name', None)
        return sorted(candidates.values(), key=lambda r: (r['name'].casefold(), r['studentId'].casefold()))

    @app.post('/api/admin/drives/<drive_id>/final-shortlist')
    @admin
    def final_shortlist(drive_id):
        data = request.get_json()
        require(isinstance(data, dict) and isinstance(data.get('studentIds'), list) and len(data['studentIds']) <= 2000,
                'Invalid shortlist request.')
        value = data.get('value', True)
        require(type(value) is bool, 'Shortlisting must be true or false.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None, 'Drive not found.')
            rows = {r['studentKey']: r for r in final_score_rows(tx, drive_id)}
            changed = 0
            for raw in data['studentIds']:
                student = clean_text(raw, 60, 'student ID')
                key = student_key(student)
                require(key in rows, 'A selected candidate is not in this drive.')
                decision = drive.setdefault('finalDecisions', {}).setdefault(key, {})
                if decision.get('interviewShortlisted') is not value:
                    decision['interviewShortlisted'] = value
                    changed += 1
                decision['interviewManual'] = False
                decision['interviewRemarks'] = ''
            return jsonify(changed=changed, serverNow=now_ms())

    @app.get('/api/admin/drives/<drive_id>/candidate-responses')
    @admin
    def candidate_responses(drive_id):
        student = clean_text(request.args.get('studentId', ''), 60, 'student ID')
        require(bool(student), 'Student ID is required.')
        key = student_key(student)
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None, 'Drive not found.')
            tests = {kind: next((test for test in tx.state['sessions']
                                 if test['driveId'] == drive_id and test['type'] == kind), None)
                     for kind in ('aptitude', 'language', 'prompt')}
            titles = {'aptitude': 'Aptitude', 'gd': 'Group Discussion (GD)',
                      'language': 'Language & Reasoning', 'prompt': 'Prompt Engineering'}
            maximums = {'aptitude': 30, 'gd': 20, 'language': 25, 'prompt': 10}
            sections = []
            names = []
            for kind in ('aptitude', 'gd', 'language', 'prompt'):
                if kind == 'gd':
                    gd = drive.get('gd') or {}
                    person = (gd.get('candidates') or {}).get(key)
                    if person:
                        score_data = (gd.get('scores') or {}).get(key) or {}
                        pool = next((p for p in gd.get('pools', []) if key in p.get('members', [])), {})
                        sections.append({
                            'type': 'gd', 'title': titles[kind], 'maxScore': maximums[kind],
                            'hasAttempt': True, 'score': score_data.get('total'),
                            'evaluator': pool.get('evaluatorName', ''),
                            'recommend': bool(score_data.get('shortlisted')),
                            'criteria': [
                                {'label': 'Initiative', 'score': score_data.get('initiative')},
                                {'label': 'Content', 'score': score_data.get('content')},
                                {'label': 'Communication', 'score': score_data.get('communication')},
                                {'label': 'Body Language', 'score': score_data.get('bodyLanguage')}
                            ]
                        })
                        names.append(person.get('name', ''))
                    else:
                        sections.append({'type': kind, 'title': titles[kind], 'maxScore': maximums[kind], 'hasAttempt': False})
                    continue
                test = tests[kind]
                if not test:
                    sections.append({'type': kind, 'title': titles[kind], 'maxScore': maximums[kind], 'hasAttempt': False})
                    continue
                attempts = [a for a in tx.all_attempts(test['id'])
                            if student_key(a.get('studentId', '')) == key]
                if not attempts:
                    sections.append({'type': kind, 'title': titles[kind], 'maxScore': maximums[kind], 'hasAttempt': False})
                    continue
                attempt = max(attempts, key=lambda a: a.get('startedAt', 0))
                names.append(attempt.get('name', ''))
                if kind == 'prompt':
                    sections.append({'type': kind, 'title': titles[kind], 'maxScore': maximums[kind],
                                     'hasAttempt': True, 'submittedAt': attempt.get('submittedAt'),
                                     'response': '\n\n'.join(
                                         (q.get('text', '') + '\n\n' if len(test['questions']) > 1 else '') +
                                         str(attempt.get('responses', {}).get(q['id']) or '')
                                         for q in test['questions']).strip() or attempt.get('answer', ''),
                                     'score': attempt.get('promptScore')})
                    continue
                graded = grade(test['questions'], attempt)
                result_by_id = {item['id']: item for item in graded['results']}
                order = attempt.get('questionOrder') or [q['id'] for q in test['questions']]
                question_map = {q['id']: q for q in test['questions']}
                ordered = [question_map[qid] for qid in order if qid in question_map]
                ordered.extend(q for q in test['questions'] if q['id'] not in set(order))
                option_order = attempt.get('optionOrder', {})
                questions = []
                for q in ordered:
                    response = attempt.get('responses', {}).get(q['id'])
                    options = q.get('options', [])
                    result = result_by_id.get(q['id'], {})
                    if q.get('options'):
                        permutation = option_order.get(q['id']) or list(range(len(options)))
                        selected = response if isinstance(response, list) else ([response] if type(response) is int else [])
                        correct_indices = q.get('key', [])
                        choices = [{'text': options[original], 'selected': original in selected,
                                    'correct': original in correct_indices}
                                   for display_index, original in enumerate(permutation)
                                   if type(original) is int and 0 <= original < len(options)]
                        answer = [choice['text'] for choice in choices if choice['selected']]
                        correct = [options[i] for i in correct_indices if type(i) is int and 0 <= i < len(options)]
                    elif q['type'] == 'matching':
                        pairs = q.get('pairs', [])
                        answer = [f"{pair.get('left', '')} → {pairs[index].get('right', '')}"
                                  for pair, index in zip(pairs, response or [])
                                  if type(index) is int and 0 <= index < len(pairs)]
                        correct = [f"{pair.get('left', '')} → {pair.get('right', '')}" for pair in pairs]
                        choices = []
                    else:
                        answer = response or ''
                        correct = q.get('referenceAnswer') or 'No fixed answer'
                        choices = []
                    questions.append({'text': q.get('text', ''), 'context': q.get('context', ''),
                                      'table': q.get('table', []), 'type': q.get('type'),
                                      'choices': choices, 'response': answer, 'correctAnswer': correct,
                                      'score': result.get('score'), 'status': result.get('status')})
                sections.append({'type': kind, 'title': titles[kind], 'maxScore': maximums[kind],
                                 'hasAttempt': True, 'submittedAt': attempt.get('submittedAt'),
                                 'score': graded['total'], 'questions': questions})
            name = next((candidate_name for candidate_name in names if candidate_name), '')
        return jsonify(driveName=drive.get('name', ''), name=name, studentId=student,
                       sections=sections, serverNow=now_ms())

    @app.post('/api/admin/drives/<drive_id>/candidate-decision')
    @admin
    def candidate_decision(drive_id):
        data = request.get_json()
        require(isinstance(data, dict), 'Invalid decision request.')
        student = clean_text(data.get('studentId'), 60, 'student ID')
        require(any(field in data for field in ('outcome', 'interviewShortlisted', 'interviewRemarks', 'interviewManual', 'interviewer')),
                'Provide a candidate decision.')
        if 'outcome' in data:
            require(data['outcome'] in {'none', 'selected', 'on-hold', 'rejected'}, 'Choose a valid candidate status.')
        if 'interviewShortlisted' in data:
            require(type(data['interviewShortlisted']) is bool, 'Choose Yes or No for interview shortlisting.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None, 'Drive not found.')
            rows = {r['studentKey'] for r in final_score_rows(tx, drive_id)}
            key = student_key(student)
            require(key in rows, 'Candidate is not in this drive.')
            decision = drive.setdefault('finalDecisions', {}).setdefault(key, {})
            if 'outcome' in data:
                decision['outcome'] = None if data['outcome'] == 'none' else data['outcome']
            if 'interviewer' in data:
                interviewer = clean_text(data['interviewer'], 100, 'interviewer name')
                members = drive.get('members', [])
                if interviewer and members:
                    canonical = next((member for member in members if student_key(member) == student_key(interviewer)), None)
                    require(canonical is not None, 'Choose an interviewer from this drive’s members.')
                    interviewer = canonical
                if decision.get('interviewManual') and decision.get('interviewShortlisted'):
                    require(bool(decision.get('interviewRemarks', '').strip()), 'Add remarks before assigning an interviewer to a manually shortlisted candidate.')
                decision['interviewer'] = interviewer
            if 'interviewShortlisted' in data or 'interviewRemarks' in data:
                value = data.get('interviewShortlisted', decision.get('interviewShortlisted', False))
                remarks = clean_text(data.get('interviewRemarks', decision.get('interviewRemarks', '')), 2000, 'interview remarks')
                manual = value and (data.get('interviewManual') is True or ('interviewShortlisted' in data and data.get('interviewManual') is not False) or decision.get('interviewManual') and data.get('interviewManual') is not False)
                require(not manual or bool(remarks), 'Remarks are required when manually shortlisting a candidate for interview.')
                decision.update(interviewShortlisted=value, interviewManual=bool(manual), interviewRemarks=remarks if value else '')
            return jsonify(outcome=decision.get('outcome') or 'none', decision=decision, serverNow=now_ms())

    @app.post('/api/admin/drives/<drive_id>/final/export')
    @admin
    def export_final_scores(drive_id):
        data = request.get_json()
        require(isinstance(data, dict) and isinstance(data.get('studentIds'), list) and len(data['studentIds']) <= 2000,
                'Invalid export request.')
        with store.transaction() as tx:
            drive = next((d for d in tx.state['drives'] if d['id'] == drive_id), None)
            require(drive is not None, 'Drive not found.')
            selected = {student_key(clean_text(value, 60, 'student ID')) for value in data['studentIds']}
            rows = [row for row in final_score_rows(tx, drive_id) if row['studentKey'] in selected]
            require(rows, 'No candidates match the current filters.')
            mode = data.get('mode', 'all')
            require(mode in {'all', 'interview'}, 'Invalid export option.')
            if mode == 'interview':
                rows = [row for row in rows if row.get('interviewShortlisted')]
            require(rows, 'No interview-shortlisted candidates in the current view.')
            payload = final_excel(drive, rows)
            name = filename(drive['name'] + ('_interview_shortlist' if mode == 'interview' else '_final_scores')) + '.xlsx'
        return send_file(io.BytesIO(payload), as_attachment=True, download_name=name,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.post('/api/admin/tests/<test_id>/marks')
    @admin
    def mark(test_id):
        data = request.get_json()
        require(isinstance(data, dict), 'Invalid request.')
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            require(test['type'] != 'prompt', 'Prompt Engineering is evaluated separately.')
            a = next((a for a in tx.all_attempts(test_id) if a['id'] == data.get('attemptId')), None)
            q = next((q for q in test['questions'] if q['id'] == data.get('questionId')), None)
            require(a and q and a.get('submittedAt'), 'Submitted candidate/question not found.')
            value = data.get('score')
            if value is None:
                a['manualMarks'].pop(q['id'], None)
            else:
                require(type(value) in (int, float) and math.isfinite(value) and q['wrongScore'] <= value <= q['correctScore'], 'Marks are outside the question range.')
                a['manualMarks'][q['id']] = value
            a['shortlisted'] = False
            return jsonify(grade=grade(test['questions'], a))

    @app.get('/api/admin/tests/<test_id>/excel')
    @admin
    def export_excel(test_id):
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            require(test['type'] != 'prompt', 'Use the Word document export for Prompt Engineering.')
            tx.expire(test_id)
            rows = results_rows(test, tx.all_attempts(test_id), *filter_numbers(request.args), filters=request.args, drive_name=next(d['name'] for d in tx.state['drives'] if d['id'] == test['driveId']))
            require(rows, 'No candidates match the filters.')
            drive = next(d for d in tx.state['drives'] if d['id'] == test['driveId'])
            payload = excel(drive, test, rows)
            name = filename(drive['name'] + '_' + test['type'] + '_results') + '.xlsx'
        return send_file(io.BytesIO(payload), as_attachment=True, download_name=name,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.get('/api/admin/tests/<test_id>/prompt-excel')
    @admin
    def export_prompt_workbook(test_id):
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            tx.expire(test_id)
            require(test['type'] == 'prompt', 'This is not a Prompt Engineering test.')
            candidates = tx.all_attempts(test_id)
            require(candidates and all(a.get('submittedAt') for a in candidates), 'Responses are not complete.')
            drive = next(d for d in tx.state['drives'] if d['id'] == test['driveId'])
            payload = prompt_workbook(drive, test, candidates)
            name = filename(drive['name'] + '_Prompt Responses') + '.xlsx'
        return send_file(io.BytesIO(payload), as_attachment=True, download_name=name,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.get('/api/admin/tests/<test_id>/word')
    @admin
    def export_word(test_id):
        with store.transaction() as tx:
            test = test_required(tx, test_id)
            tx.expire(test_id)
            require(test['type'] == 'prompt', 'This is not a Prompt Engineering test.')
            candidates = tx.all_attempts(test_id)
            require(candidates and all(a.get('submittedAt') for a in candidates), 'Responses are not complete.')
            drive = next(d for d in tx.state['drives'] if d['id'] == test['driveId'])
            payload = words(drive, test, candidates)
            name = filename(test['name'] + '_responses') + '.zip'
        return send_file(io.BytesIO(payload), as_attachment=True, download_name=name, mimetype='application/zip')

    if ready() and not local:
        def expiry_worker():
            while True:
                try:
                    with store.transaction() as tx:
                        tx.expire()
                except Exception as error:
                    app.logger.error('Expiry worker could not reach storage: %s', type(error).__name__)
                time.sleep(5)
        threading.Thread(target=expiry_worker, name='assessment-expiry', daemon=True).start()
    return app


application = create_app()

